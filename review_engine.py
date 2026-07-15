"""규칙 기반 자가검토 엔진.

파싱된 조서(ParsedDocument)를 입력받아 형식·완전성, 계산검증, 증빙·절차
규칙을 적용하고 리뷰노트를 생성합니다. (AI·RAG 연동 전 단계)

근거(basis)는 일반적인 회계·감사기준 조항만 인용하며,
검증되지 않은 감리지적 번호 등은 만들어 내지 않습니다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Any, Callable

import pandas as pd
from openpyxl.utils import get_column_letter as _col_letter

import config as app_config
import fss_focus
import enforcement_review
import knowledge_base as kb
import qc_review
import sheet_code_registry as scr
from parser import ParsedDocument, to_number


@dataclass
class Materiality:
    """조서에서 추출·입력한 중요성 금액."""

    te: float | None = None
    pm: float | None = None
    ctt: float | None = None

    @property
    def effective(self) -> float | None:
        return self.pm or self.te

    def min_calc_threshold(self) -> float:
        base = 1000.0
        eff = self.effective
        if eff and eff > 0:
            return max(base, eff * 0.01)
        return base

# 합계 성격의 행/열을 식별하는 표현
_TOTAL_RE = re.compile(r"합\s*계|총\s*계|^\s*계\s*$|total", re.IGNORECASE)
_SUBTOTAL_RE = re.compile(r"소\s*계|중\s*계|sub", re.IGNORECASE)

_IMPORTANCE_BY_CATEGORY = {
    "계산검증": "하",
    "증빙·절차": "중",
    "절차누락": "상",
    "형식·완전성": "하",
    "중점감리": "상",
    "주석검증": "상",
    "중요성": "상",
    "QC·품질관리": "중",
    "개선제안": "하",
    "검토요청": "중",
}

_MAT_RE = re.compile(
    r"(?:중요성\s*금액|Overall\s*Materiality|TE)\s*[:：]?\s*([\d,]+(?:\.\d+)?)",
    re.IGNORECASE,
)
_PM_RE = re.compile(
    r"(?:수행\s*중요성|Performance\s*Materiality|PM)\s*[:：]?\s*([\d,]+(?:\.\d+)?)",
    re.IGNORECASE,
)
_CTT_RE = re.compile(
    r"(?:명백히\s*사소한|CTT|SAD)\s*[:：]?\s*([\d,]+(?:\.\d+)?)",
    re.IGNORECASE,
)


def extract_engagement(doc: ParsedDocument) -> dict[str, Any]:
    """조서 텍스트에서 기본정보를 추출(가능한 범위)."""
    text = doc.text

    company = _extract_company_from_lead_sheets(doc)
    if not company:
        company = _extract_company(text, doc.file_name)

    year = _extract_audit_year_from_lead_sheets(doc)
    if not year:
        year = _extract_audit_year(text, doc.file_name)

    preparer = _extract_person(text, ["작성자", "Preparer", "Prepared\\s*by"])
    reviewer = _extract_person(
        text, ["검토자", "심리자", "Reviewer", "Reviewed\\s*by"]
    )

    listed = detect_is_listed(doc, company or doc.file_name)

    if "K-IFRS" in text or "한국채택국제회계기준" in text:
        accounting = "K-IFRS"
    elif "일반기업회계기준" in text:
        accounting = "일반기업회계기준"
    elif not listed:
        accounting = "일반기업회계기준"
    else:
        accounting = "K-IFRS"

    accounts = document_accounts(doc)
    mat = extract_materiality(doc)

    workpapers = doc.source_files if doc.source_files else [doc.file_name]

    return {
        "company_name": (
            company or (doc.file_name.rsplit(".", 1)[0] if "." in doc.file_name else doc.file_name)
        ).strip(),
        "audit_year": (year or "확인 필요").strip(),
        "accounting_standard": accounting,
        "audit_standard": "회계감사기준(KSA)",
        "related_workpaper": ", ".join(workpapers),
        "related_account": " / ".join(accounts) if accounts else "확인 필요",
        "preparer": (preparer or "확인 필요").strip(),
        "reviewer": (reviewer or "").strip(),
        "review_status": "심리 제출 전 자가검토 진행중",
        "materiality_te": mat.te,
        "materiality_pm": mat.pm,
        "materiality_ctt": mat.ctt,
        "materiality_display": _format_materiality(mat),
        "is_listed": listed,
    }


_CORP_PATTERNS = [
    r"㈜\s*([가-힣A-Za-z0-9]{2,20})",
    r"\(\s*주\s*\)\s*([가-힣A-Za-z0-9]{2,20})",
    r"주식회사\s*([가-힣A-Za-z0-9]{2,20})",
    r"([가-힣A-Za-z0-9]{2,20})\s*주식회사",
]


def _extract_person(text: str, labels: list[str]) -> str:
    """작성자/검토자 이름을 추출. 한글·영문(Preparer/Reviewer) 라벨 모두 지원.

    라벨 뒤의 한글 이름(예: 김재형, 김○○)을 우선 잡고, 날짜·직함은 제외한다.
    """
    for lab in labels:
        # 라벨 [:] 뒤 한글 이름 (회계사/CPA 직함은 이름에서 분리)
        m = re.search(
            lab + r"\s*[:：]?\s*([가-힣○]{2,10})(?=\s|$|회계사|CPA|,|\d)",
            text,
            re.IGNORECASE,
        )
        if m:
            name = m.group(1).strip()
            if name and not name.isdigit():
                return name
    return ""


def _extract_company(text: str, fallback: str) -> str:
    """회사명을 추출. 법인 형태(㈜/(주)/주식회사) 토큰을 우선한다."""
    head = "\n".join(text.split("\n")[:40])
    name = _extract_company_from_text_head(head)
    if name:
        return name
    return fallback


_LEAD_HINT_RE = re.compile(r"lead|리드", re.I)


def _extract_company_from_text_head(head: str) -> str | None:
    """시트 상단(약 25행)에서 회사명 추출."""
    for pat in _CORP_PATTERNS:
        m = re.search(pat, head)
        if m:
            raw = m.group(0)
            if m.lastindex:
                inner = m.group(1)
                prefix = "㈜" if raw.startswith("㈜") else "(주)" if "(주" in raw else ""
                if prefix:
                    return prefix + re.sub(r"\s+", "", inner)
            return re.sub(r"\s+", "", raw)
    m = re.search(
        r"(?:회사\s*명|상호|피감사\s*법인|법인\s*명)\s*[:：]\s*([^\n]{2,40})",
        head,
    )
    if m:
        name = _clean_company(m.group(1))
        if name:
            return name
    return None


def _is_lead_like_sheet(t: pd.DataFrame) -> bool:
    label = f"{t.attrs.get('source', '')} {t.attrs.get('title', '')}"
    if _LEAD_HINT_RE.search(label):
        return True
    source = str(t.attrs.get("source", "")).strip()
    code_part = source.split("/")[-1] if "/" in source else source
    return _account_from_code(code_part.split()[0] if code_part else "") is not None


def _extract_company_from_lead_sheets(doc: ParsedDocument) -> str | None:
    """Lead·계정별 조서 상단에서 회사명 추출 — 최빈값(다수결) 채택."""
    from collections import Counter

    weights: Counter[str] = Counter()
    for t in doc.tables:
        name = _extract_company_from_text_head(_sheet_text(t)[:3000])
        if not name:
            continue
        w = 3 if _LEAD_HINT_RE.search(f"{t.attrs.get('source', '')} {t.attrs.get('title', '')}") else 1
        if _is_lead_like_sheet(t):
            weights[name] += w
    if not weights:
        return None
    return weights.most_common(1)[0][0]


def _clean_company(value: str) -> str:
    """라벨 뒤 텍스트에서 회사명만 정제 (숫자·잡음 제거)."""
    value = value.strip()
    # 숫자가 2개 이상 연속되면 그 앞까지만 (계정코드·금액 방지)
    value = re.split(r"\s*\d{2,}", value)[0].strip()
    # 다른 라벨이 이어지면 절단
    value = re.split(r"\s{2,}|[:：]", value)[0].strip()
    # 이름이 숫자뿐이거나 너무 짧으면 무효
    if not value or value.isdigit() or len(value) < 2:
        return ""
    return value


def _norm_year(y: str) -> str:
    """2자리 연도(24)를 4자리(2024)로 정규화."""
    y = y.strip()
    return f"20{y}" if len(y) == 2 else y


_PERIOD_A_RE = re.compile(
    r"A\s*[:：]\s*(20\d{2})\s*[-./]\s*12\s*[-./]\s*31",
    re.I,
)
_PERIOD_LABEL_RE = re.compile(
    r"(?:회계년도|회계연도|결산연도|사업연도|"
    r"결산일|보고기간\s*종료일|결산기준일)\s*[:：]?\s*(20\d{2})",
    re.I,
)


def _extract_audit_year_from_sheet_head(head: str) -> str | None:
    """Lead·계정별 조서 좌측 상단의 당기 결산연도 추출."""
    for pat in (_PERIOD_A_RE, _PERIOD_LABEL_RE):
        m = pat.search(head)
        if m:
            y = m.group(1)
            if 2000 <= int(y) <= 2035:
                return y
    return None


def _extract_audit_year_from_lead_sheets(doc: ParsedDocument) -> str | None:
    """Lead·계정별 조서 상단 `A: YYYY-12-31` 등에서 감사연도 추출 — 최빈값 채택."""
    from collections import Counter

    weights: Counter[str] = Counter()
    for t in doc.tables:
        head = _sheet_text(t)[:1200]
        year = _extract_audit_year_from_sheet_head(head)
        if not year:
            continue
        label = f"{t.attrs.get('source', '')} {t.attrs.get('title', '')}"
        if _LEAD_HINT_RE.search(label):
            w = 3
        elif _is_lead_like_sheet(t):
            w = 2
        else:
            w = 1
        weights[year] += w
    if not weights:
        return None
    return weights.most_common(1)[0][0]


def _extract_audit_year(text: str, file_name: str) -> str | None:
    """감사(대상)연도를 결산일·연도말 날짜 기준으로 추출.

    문서 내 임의의 20xx 숫자(계정코드 등)를 잘못 잡지 않도록,
    결산일/회계연도말(12/31) 날짜를 우선 사용한다.
    전기 비교용 「기준일」은 당기 연도로 오인하지 않는다.
    """
    import collections

    # 0) 당기 결산일 A: 표기 (Lead 추출 실패 시 본문 폴백)
    m = _PERIOD_A_RE.search(text)
    if m:
        return m.group(1)

    # 1) 결산일·보고기간종료일 등 라벨 + 연도 (기준일 제외 — 전기 비교일 오탐 방지)
    m = _PERIOD_LABEL_RE.search(text)
    if m:
        return m.group(1)

    # 2) 회계연도말 날짜 패턴 (YYYY-12-31 / 12.31.24 등)
    m = re.search(r"(20\d{2})\s*[-./]\s*12\s*[-./]\s*31", text)
    if m:
        return m.group(1)
    m = re.search(r"\b12\s*[-./]\s*31\s*[-./]\s*(\d{2})\b", text)
    if m:
        return _norm_year(m.group(1))

    # 3) FY 표기 (본문·파일명)
    m = re.search(r"(?:FY|fy|회계연도)\s*[:\-]?\s*(20\d{2})", f"{text} {file_name}")
    if m:
        return m.group(1)

    # 4) 문서 내 완전한 날짜(YYYY-MM-DD 등)의 연도 최빈값
    yrs = re.findall(r"(20\d{2})\s*[-./]\s*\d{1,2}\s*[-./]\s*\d{1,2}", text)
    yrs = [y for y in yrs if 2000 <= int(y) <= 2035]
    if yrs:
        return collections.Counter(yrs).most_common(1)[0][0]

    # 5) 파일명의 4자리 연도
    m = re.search(r"(20\d{2})", file_name)
    if m:
        return m.group(1)
    return None


def _parse_amount(raw: str) -> float | None:
    try:
        return float(raw.replace(",", "").strip())
    except ValueError:
        return None


def extract_materiality(doc: ParsedDocument) -> Materiality:
    """계획·총괄 조서에서 중요성 금액 추출."""
    text = doc.text
    te = pm = ctt = None
    m = _MAT_RE.search(text)
    if m:
        te = _parse_amount(m.group(1))
    m = _PM_RE.search(text)
    if m:
        pm = _parse_amount(m.group(1))
    m = _CTT_RE.search(text)
    if m:
        ctt = _parse_amount(m.group(1))
    return Materiality(te=te, pm=pm, ctt=ctt)


def detect_is_listed(doc: ParsedDocument, company: str = "") -> bool:
    """상장법인 여부 추정 (4대 중점항목 매핑용)."""
    hay = f"{doc.text} {company} {doc.file_name}"
    if re.search(r"주권\s*상장|코스피|코스닥|KRX|유가증권\s*시장", hay):
        return True
    if re.search(r"비\s*상장|비상장", hay):
        return False
    return True  # 미확인 시 상장 기준(금감원) 적용


def _format_materiality(mat: Materiality) -> str:
    parts: list[str] = []
    if mat.te:
        parts.append(f"TE {mat.te:,.0f}")
    if mat.pm:
        parts.append(f"PM {mat.pm:,.0f}")
    if mat.ctt:
        parts.append(f"CTT {mat.ctt:,.0f}")
    return " / ".join(parts) if parts else "미확인"


def run_review(
    doc: ParsedDocument,
    include_minor: bool = False,
    materiality: Materiality | None = None,
    is_listed: bool | None = None,
    engagement: dict[str, Any] | None = None,
    progress: Callable[[str], None] | None = None,
) -> list[dict[str, Any]]:
    """규칙을 적용해 리뷰노트 목록을 생성.

    본문 조서의 작성자·검토자·작성일 미기재 등 형식적 사항은 중요하지 않아
    기본적으로 지적하지 않는다(include_minor=True 로 두면 포함).
    핵심은 ①증빙·외부조회 등 중요 감사절차 누락 ②표준조서 대비 절차 누락
    ③명백한 계산오류 이다.
    """
    def _prog(msg: str) -> None:
        if progress:
            progress(msg)

    notes: list[dict[str, Any]] = []
    mat = materiality or extract_materiality(doc)
    listed = is_listed if is_listed is not None else detect_is_listed(doc)
    scr.set_mapping_variant(listed)

    if include_minor:
        _prog("형식·완전성 점검…")
        # 형식·완전성, 중요성 문서화, 별첨 대조, 계정 매핑 진단 등은
        # 조서 자체의 중대한 흠결이 아니므로 QRM 제출용 상세 모드에서만 지적한다.
        notes += _check_materiality_missing(doc, mat)
        notes += _check_completeness(doc)
        notes += _check_evidence(doc)
        notes += _check_account_mapping(doc)
    _prog("표준절차·계정별 점검…")
    notes += _check_procedures(doc, mat, listed)
    eng = engagement or extract_engagement(doc)
    _prog("4대 중점·감리지적 체크리스트…")
    notes += fss_focus.run_focus_review(
        doc,
        {"audit_year": eng.get("audit_year") or _extract_audit_year(doc.text, doc.file_name)},
        is_listed=listed,
    )
    notes += enforcement_review.run_enforcement_checklist_review(doc, is_listed=listed)
    _prog("합계·주석 대사 점검…")
    notes += _check_calculations(doc, mat)
    notes += _check_footnotes(doc, mat, include_minor=include_minor)
    notes += _check_note_references(doc, mat)
    notes += _check_contingency_confirmations(doc)
    notes += _check_cross_fs_totals(doc, mat)
    if include_minor:
        notes += _check_prior_year_analysis(doc)
    _prog("품질관리(QC) 점검…")
    notes += qc_review.run_qc_checks(doc, eng)

    if not include_minor:
        # 중요성·맥락상 경미한 항목(중요도 '하')은 패스 — 정말 중요한 누락만 보고.
        notes = [
            n for n in notes
            if n.get("importance") != "하"
            or n.get("focus_protected")
            or n.get("enforcement_protected")
            or n.get("category") in ("중점감리", "감리지적체크")
        ]
    from notes_pipeline import filter_cross_sheet_procedure_notes

    notes = filter_cross_sheet_procedure_notes(doc, notes)
    return _assign_ids(notes)


def _doc_sheet(doc: ParsedDocument) -> tuple[str, str]:
    """문서 전반을 가리키는 (조서번호, 계정과목)."""
    if doc.tables:
        t = doc.tables[0]
        return t.attrs.get("source", ""), t.attrs.get("title", "")
    if doc.sheet_names:
        return doc.sheet_names[0], ""
    return "", ""


# 짧은 계정 동의어가 다른 계정 용어의 일부로 오인되는 것을 막기 위한 예외.
# 예: 재고자산 '상품'이 '금융상품'(현금·예금 계정)에 오탐되지 않도록 한다.
_SYNONYM_EXCLUSIONS: dict[str, tuple[str, ...]] = {
    "상품": ("금융상품", "상품권"),
    "제품": ("금융상품",),
    "매출": ("매출채권",),
    "현금": ("현금흐름",),
}


def _synonym_pos(text: str, synonym: str) -> int | None:
    """동의어가 처음 등장하는 위치(index). 오탐 유발 복합어는 위치를 보존한 채 무효화."""
    excl = _SYNONYM_EXCLUSIONS.get(synonym)
    if excl:
        for e in excl:
            text = text.replace(e, " " * len(e))  # 길이 유지로 위치 보존
    idx = text.find(synonym)
    return idx if idx >= 0 else None


def _sheet_text(table: pd.DataFrame) -> str:
    """시트(표)의 셀 내용을 문자열로 합친다 (해당 시트 본문 검색용)."""
    cached = table.attrs.get("_cached_sheet_text")
    if isinstance(cached, str):
        return cached
    parts: list[str] = []
    for row in table.itertuples(index=False, name=None):
        parts.append(" ".join("" if c is None else str(c) for c in row))
    text = "\n".join(parts)
    table.attrs["_cached_sheet_text"] = text
    return text


# 계정과목 분류 체계 (표준명, 동의어). 조서 시트를 '주(主)계정'으로 매핑하는 데 사용.
#   · PROCEDURE_RULES 에 있는 계정(매출채권·재고자산·매입채무·차입금·현금및예금·매출)만
#     절차누락 점검 대상이다.
#   · 그 외 계정(우발부채·영업외손익·법인세 등)도 등록해, 해당 시트를 '선점'하게 하여
#     절차 규칙이 오탐으로 그 시트를 가로채지 않도록 한다.
ACCOUNT_TAXONOMY: list[tuple[str, tuple[str, ...]]] = [
    ("매출채권", ("매출채권", "외상매출금", "받을어음")),
    ("유가증권", ("유가증권", "지분법", "매도가능증권")),
    ("파생상품", ("파생상품", "파생금융")),
    ("기타유동자산", ("기타유동자산", "선급금", "선급비용", "미수금")),
    ("기타유동부채", ("기타유동부채", "미지급금", "미지급비용", "예수금", "선수금", "가수금", "유동성장기부채")),
    ("기타비유동자산", ("기타비유동자산",)),
    ("기타비유동부채", ("기타비유동부채", "기타의 비유동부채")),
    ("대손충당금", ("대손충당금", "기대신용손실", "손실충당금")),
    ("현금및예금", ("현금및현금성자산", "현금성자산", "장단기금융상품", "현예금", "현금및예금", "예금", "장기예금", "장기성예금")),
    ("재고자산", ("재고자산", "상품", "제품", "원재료", "재공품", "미착품", "저장품", "재고자산실사입회")),
    ("유형자산", ("유형자산", "토지", "건물", "구축물", "기계장치", "차량운반구", "비품", "감가상각누계액")),
    ("투자부동산", ("투자부동산",)),
    ("무형자산", ("무형자산", "영업권", "개발비", "산업재산권", "소프트웨어")),
    ("투자자산", (
        "투자자산", "관계기업투자", "종속기업투자", "매도가능증권",
        "종속기업", "관계기업", "공동기업", "지분법적용투자주식", "공동기업투자주식",
    )),
    ("매입채무", ("매입채무", "외상매입금", "지급어음")),
    ("차입금", ("차입금", "사채", "단기차입", "장기차입", "장단기차입금")),
    ("충당부채", ("충당부채", "제품보증충당", "제충당부채")),
    ("퇴직급여", ("퇴직급여충당", "퇴직급여", "확정급여", "퇴직급여부채", "급여")),
    ("우발부채·약정", ("우발부채", "약정사항", "지급보증", "계류중인 소송", "우발부채 및 약정사항")),
    ("부외부채", ("부외부채", "부외", "팩토링", "일광", "SPE", "우선수익")),
    ("자본", ("자본금", "자본잉여금", "이익잉여금", "자본조정")),
    ("매출", ("매출액", "매출", "수익", "매출 및 기타수익")),
    ("매출원가", ("매출원가", "제조원가", "매출원가 및 제조원가")),
    ("판매관리비", ("판매비와관리비", "판매비와 관리비", "판관비")),
    ("영업외손익", ("영업외손익", "영업외수익", "영업외비용", "기타수익", "기타비용")),
    ("법인세", ("이연법인세", "미지급법인세", "법인세비용", "법인세")),
    ("리스", ("리스", "사용권자산", "리스부채")),
    ("자산손상", ("자산손상", "손상차손")),
    ("매각예정자산", ("매각예정", "매각예정비유동자산", "중단사업")),
    ("중단사업", ("중단사업", "매각예정비유동자산과 중단사업")),
    ("공동약정", ("공동약정",)),
    ("조인트벤처", ("조인트벤처", "Joint venture")),
]


# 조서 인덱스 코드 → 계정과목 (한울 FY2026 표준).
# 출처: data/sheet_code_map_fy2026.json (상장 IFRS / 비상장 목차)
# 비상장·상장 차이(F·P·U·JV 등)는 sheet_code_registry.set_mapping_variant() 로 분기.
SHEET_CODE_MAP: dict[str, str] = scr.build_sheet_code_map(variant="ifrs_listed")


def _account_from_code(label: str) -> str | None:
    """시트명(조서번호)의 영문 코드부로 계정과목을 판정."""
    return scr.account_from_label(label)


_JUNK_LABEL_RE = re.compile(
    r"^(?:\d{4}-\d{2}-\d{2}(?:\s+\d{2}:\d{2}:\d{2})?|\d{2}:\d{2}:\d{2})$"
)
_QC_TOPIC_TITLES = frozenset({"계속기업", "독립성", "QC·품질관리", "QC·전체조서"})


def _sanitize_account_label(label: str) -> str:
    """결산일·시간 등 계정명이 아닌 값을 제거."""
    s = str(label or "").strip()
    if not s or _JUNK_LABEL_RE.match(s):
        return ""
    s = re.sub(r"^\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}:\d{2}\s*", "", s).strip()
    return "" if _JUNK_LABEL_RE.match(s) else s


def note_display_account(note: dict[str, Any]) -> str:
    """UI·통합 노트에 표시할 계정과목(조서 코드 우선, QC 주제명 제외)."""
    acct = note_account(note)
    if acct:
        return acct
    st = _sanitize_account_label(str(note.get("sheet_title") or ""))
    if st and st not in _QC_TOPIC_TITLES:
        return st
    return ""


def table_account(t: pd.DataFrame) -> str | None:
    """시트의 주계정 판정 — ①조서번호 코드 ②제목 ③본문 최초 등장 순."""
    acct = _account_from_code(str(t.attrs.get("source", "")).strip())
    if acct:
        return acct
    title = _sanitize_account_label(str(t.attrs.get("title", "")).strip())
    acct = _account_from_code(title)
    if acct:
        return acct
    return _primary_account(title, _sheet_text(t))


def _primary_account(title: str, body_text: str) -> str | None:
    """시트의 '주(主)계정'을 판정한다.

    리드시트는 상단에 그 시트의 계정명이 먼저 나오므로,
    ① 제목에서 계정명을 찾고 ② 없으면 본문에서 '가장 먼저' 등장하는 계정명을 택한다.
    빈도(합계 덤프에서 가장 많이 나온 계정)가 아니라 위치를 기준으로 하여,
    시트 하단의 재무제표 덤프에 휩쓸리지 않게 한다.
    """
    for src in (title or "", body_text or ""):
        best_name: str | None = None
        best_pos: int | None = None
        for name, syns in ACCOUNT_TAXONOMY:
            for s in syns:
                pos = _synonym_pos(src, s)
                if pos is not None and (best_pos is None or pos < best_pos):
                    best_pos, best_name = pos, name
        if best_name:
            return best_name
    return None


def note_account(note: dict[str, Any]) -> str | None:
    """리뷰노트가 다루는 계정과목 판정 (조서번호 코드 → 제목·지적문 순)."""
    acct = _account_from_code(str(note.get("sheet_no", "")))
    if acct:
        return acct
    hay = f"{note.get('account', '')} {note.get('sheet_title', '')}"
    return _primary_account(hay, str(note.get("defect", "")))


# 계정별 무관 주제 — 타 계정·업종 이슈가 리뷰노트·감리사례에 섞이는 것을 차단
ACCOUNT_TOPIC_REJECT: dict[str, tuple[str, ...]] = {
    "현금및예금": (
        "공사", "건설", "장기공사", "공사계약", "공사수익", "공사원가", "수익인식",
        "매출원가", "재고자산", "재고", "실사입회", "cut-off", "cutoff", "컷오프",
        "대손", "매출채권", "채권", "퇴직", "법인세", "percentage of completion",
    ),
    "매출": ("예금", "현금성", "은행조회", "재고실사", "차입금", "사채", "대손충당"),
    "재고자산": ("예금", "현금", "공사계약", "수익인식", "차입", "대손", "은행조회"),
    "매출채권": ("재고", "공사", "예금", "차입", "법인세"),
    "차입금": ("재고", "매출인식", "수익인식", "공사", "대손"),
    "우발부채·약정": (
        "수익", "비용", "기간귀속", "매출", "cut-off", "cutoff", "컷오프",
        "재고", "공사", "현금", "예금", "대손", "법인세", "수익인식",
    ),
}
ACCOUNT_TOPIC_ALLOW: dict[str, tuple[str, ...]] = {
    "현금및예금": (
        "현금", "예금", "금융", "은행", "조회", "장기예금", "장기성", "이자", "만기",
        "금융상품", "확정", "당좌", "정기", "잔액", "deposit",
    ),
    "우발부채·약정": (
        "우발", "약정", "지급보증", "소송", "담보", "충당", "계류", "보증", "연대",
        "pf", "풋옵션", "채무인수",
    ),
}
_FOREIGN_TOPIC_PATTERNS: list[re.Pattern[str]] = [
    re.compile(r"장기\s?공사|공사\s?계약|공사\s?수익|공사\s?원가|건설\s?계약", re.I),
    re.compile(r"수익\s?인식|매출\s?인식|percentage|완성\s?기준|공사\s?기간", re.I),
    re.compile(r"재고\s?실사|cut.?off|컷오프|실사\s?입회", re.I),
    re.compile(r"대손\s?충당|기대\s?신용|ECL", re.I),
]


def topic_off_account_text(text: str, account: str | None) -> bool:
    """계정과 무관한 주제어(공사수익·재고실사 등)가 텍스트에 있는지."""
    if not account or not text:
        return False
    reject = ACCOUNT_TOPIC_REJECT.get(account)
    if not reject:
        return False
    hay_low = text.lower()
    if not any(r.lower() in hay_low if len(r) > 2 else r in text for r in reject):
        return False
    allow = ACCOUNT_TOPIC_ALLOW.get(account, ())
    if any(a.lower() in hay_low if len(a) > 2 else a in text for a in allow):
        return False
    for name, syns in ACCOUNT_TAXONOMY:
        if name == account and any(_synonym_pos(text, s) is not None for s in (name, *syns)):
            return False
    return True


def _case_title_text(source: str) -> str:
    return source.split("·")[-1].strip() if source else ""


def is_off_account_enforcement(source: str, snippet: str, note_acct: str | None) -> bool:
    """감리지적사례 — **제목**이 노트 계정·주제와 무관하면 차단 (스니펫만으로 통과 금지)."""
    if not note_acct:
        return False
    title = _case_title_text(source)
    if title and topic_off_account_text(title, note_acct):
        return True
    if title:
        reject = ACCOUNT_TOPIC_REJECT.get(note_acct, ())
        allow = ACCOUNT_TOPIC_ALLOW.get(note_acct, ())
        t_low = title.lower()
        if reject and any(r.lower() in t_low for r in reject):
            if not any(a.lower() in t_low for a in allow):
                return True
    return topic_off_account_text(f"{source} {snippet}", note_acct)


def is_off_account_text(text: str, account: str | None) -> bool:
    """근거·사례 텍스트가 해당 계정이 아닌 '다른 계정'에 관한 것인지 판정.

    예: 유형자산 노트에 매출채권 감리지적사례가 붙는 것을 차단한다.
    자기 계정이 언급되어 있으면 관련 있는 것으로 보고 통과시킨다.
    """
    if not account:
        return False
    if topic_off_account_text(text, account):
        return True
    own_syns: list[str] = [account]
    for name, syns in ACCOUNT_TAXONOMY:
        if name == account:
            own_syns.extend(syns)
            break
    if any(_synonym_pos(text, s) is not None for s in own_syns):
        return False
    masked = text
    for s in sorted(own_syns, key=len, reverse=True):
        masked = masked.replace(s, " ")
    for name, syns in ACCOUNT_TAXONOMY:
        if name == account:
            continue
        for s in (name, *syns):
            if _synonym_pos(masked, s) is not None:
                return True
    return False


def _sheet_text_for_note(doc: ParsedDocument, note: dict[str, Any]) -> str:
    target = str(note.get("sheet_no") or note.get("workpaper_ref") or "")
    if not target:
        return ""
    parts: list[str] = []
    for t in doc.tables:
        label = f"{t.attrs.get('source', '')} {t.attrs.get('title', '')}".strip()
        if _sheet_keys_match(target, label):
            parts.append(_sheet_text(t))
    return "\n".join(parts)


def note_topic_absent_from_sheet(note: dict[str, Any], doc: ParsedDocument) -> bool:
    """지적 주제가 해당 시트 본문에 없으면 타 조서 이슈로 간주."""
    hay = f"{note.get('defect', '')} {note.get('reason', '')}"
    sheet = _sheet_text_for_note(doc, note)
    if not sheet.strip():
        return False
    for pat in _FOREIGN_TOPIC_PATTERNS:
        if pat.search(hay) and not pat.search(sheet):
            return True
    return False


def is_off_account_note(note: dict[str, Any], doc: ParsedDocument | None = None) -> bool:
    """리뷰노트가 시트 계정·본문과 무관한 지적인지."""
    if note.get("category") == "QC·품질관리" and str(note.get("location") or "").strip():
        return False
    acct = note_account(note)
    hay = f"{note.get('defect', '')} {note.get('reason', '')} {note.get('to_be', '')}"
    if is_off_account_text(hay, acct):
        return True
    if doc and note_topic_absent_from_sheet(note, doc):
        return True
    return False


def sanitize_note_citations(note: dict[str, Any]) -> None:
    """노트 계정과 무관한 근거·감리지적사례 첨부를 제거."""
    acct = note_account(note)
    refs = note.get("references") or []
    kept_refs = [
        r for r in refs
        if not is_off_account_text(f"{r.get('source', '')} {r.get('snippet', '')}", acct)
    ]
    if kept_refs:
        note["references"] = kept_refs
    else:
        note.pop("references", None)
    cases = note.get("enforcement_cases") or []
    kept_cases = [
        c for c in cases
        if not is_off_account_enforcement(
            f"{c.get('subject', '')} {c.get('number', '')}",
            f"{c.get('brief', '')} {c.get('summary_line', '')}",
            acct,
        )
    ]
    if kept_cases:
        note["enforcement_cases"] = kept_cases
    else:
        note.pop("enforcement_cases", None)


_BORROWING_COLLATERAL_NOTE_RE = re.compile(
    r"담보.{0,32}(?:제공|설정|채권|자산)|차입.{0,24}담보|"
    r"collateral|pledge|근저당|질권",
    re.I,
)
_COLLATERAL_CL_KW = (
    "담보", "근저당", "질권", "담보제공", "담보설정", "담보자산", "제공자산", "설정자산",
)
_BORROWING_CTX_KW = ("차입", "대출", "차입금", "사채", "금융기관", "borrow")


def _is_contingency_sheet(t: pd.DataFrame) -> bool:
    if table_account(t) == "우발부채·약정":
        return True
    label = f"{t.attrs.get('source', '')} {t.attrs.get('title', '')}"
    code_part = str(t.attrs.get("source", "")).split("/")[-1]
    if _account_from_code(code_part.split()[0] if code_part else "") == "우발부채·약정":
        return True
    return bool(re.search(r"우발|약정|(?:^|\s|/)CL(?:\s|$|Lead|\d|리드)", label, re.I))


def doc_has_borrowing_collateral_in_contingency(doc: ParsedDocument) -> bool:
    """우발부채·약정(CL) 조서에 차입 담보·제공자산 기재가 있는지."""
    for t in doc.tables:
        if not _is_contingency_sheet(t):
            continue
        text = _sheet_text(t).lower()
        if not any(k in text for k in _COLLATERAL_CL_KW):
            continue
        if any(k in text for k in _BORROWING_CTX_KW):
            return True
        if any(k in text for k in ("제공", "설정", "한도", "금액", "장부")):
            return True
    doc_ev = _doc_evidence(doc).lower()
    if any(k in doc_ev for k in _COLLATERAL_CL_KW):
        if re.search(r"cl\s|/cl|우발|약정", doc_ev) and any(k in doc_ev for k in _BORROWING_CTX_KW):
            return True
    return False


def is_borrowing_collateral_disclosure_note(note: dict[str, Any]) -> bool:
    """차입금 조서 주석에 담보 공시 누락을 지적하는 노트."""
    hay = f"{note.get('defect', '')} {note.get('reason', '')} {note.get('to_be', '')}"
    if not _BORROWING_COLLATERAL_NOTE_RE.search(hay):
        return False
    acct = note_account(note)
    sheet = f"{note.get('sheet_no', '')} {note.get('sheet_title', '')}"
    if acct == "차입금":
        return True
    return bool(re.search(r"차입|사채|BB|DD", hay + sheet, re.I))


def soft_borrowing_collateral_memo(note: dict[str, Any]) -> dict[str, Any]:
    """차입 담보 — 우발·약정 조서 확인 수준의 경미 메모."""
    memo = dict(note)
    memo["importance"] = "중"
    memo["category"] = "증빙·절차"
    memo["defect"] = "[확인] 차입금 담보·제공 내역 — 우발부채·약정(CL) 조서 기재 여부"
    memo["reason"] = (
        "차입금 관련 담보 제공 내역은 통상 우발부채·약정사항 조서에 기재합니다. "
        "차입금(BB·DD) 조서 주석에 별도 공시가 없더라도, 우발부채·약정(CL) 조서에 "
        "차입 담보·제공자산 내역이 적정하게 기록·대사되어 있는지 확인하십시오."
    )
    memo["to_be"] = "우발부채·약정 조서에서 차입 담보·제공자산 내역을 확인·교차참조하십시오."
    memo["basis"] = "실무 관행(담보·약정사항은 CL 조서에서 관리)"
    memo["collateral_memo"] = True
    memo.pop("enforcement_cases", None)
    memo.pop("references", None)
    return memo


def document_accounts(doc: ParsedDocument) -> list[str]:
    """조서에서 실제로 다뤄진 계정과목을 시트 단위로 수집(주계정 기준)."""
    names: list[str] = []

    def add(value: str) -> None:
        value = _sanitize_account_label(value)
        if value and value not in names:
            names.append(value)

    for t in doc.tables:
        code_acct = _account_from_code(str(t.attrs.get("source", "")).strip())
        title = _sanitize_account_label(str(t.attrs.get("title", "")).strip())
        if code_acct:
            add(code_acct)  # 조서번호 코드가 가장 신뢰도 높음 (예: C → 매출채권)
        elif title:
            add(title)  # 시트가 스스로 밝힌 계정명을 그대로 표기
        else:
            acct = _primary_account("", _sheet_text(t))
            if acct:
                add(acct)
    if names:
        return names
    # 표가 없는 문서(텍스트 PDF 등)는 본문 최초 등장 계정으로 보조 판정
    acct = _primary_account("", doc.text)
    return [acct] if acct else []


# --- 규칙 ① 형식·완전성 ---
def _check_completeness(doc: ParsedDocument) -> list[dict[str, Any]]:
    text = doc.text
    notes: list[dict[str, Any]] = []

    checks = [
        ("작성일", r"작성\s*일", "조서 작성일이 기재되지 않았습니다."),
        ("검토일", r"(?:검토|심리)\s*일", "1차 검토일(심리일)이 기재되지 않았습니다."),
        ("작성자", r"작성자", "조서 작성자가 기재되지 않았습니다."),
        ("검토자", r"검토자|심리자", "검토자(심리자)가 기재되지 않았습니다."),
    ]
    sheet_no, sheet_title = _doc_sheet(doc)
    for label, pattern, msg in checks:
        if not re.search(pattern, text):
            notes.append(
                _note(
                    category="형식·완전성",
                    defect=f"{label} 미기재",
                    reason=msg,
                    basis="법인 표준조서 양식 필수 항목; 품질관리 매뉴얼 조서완결성 체크리스트",
                    to_be=f"{label} 항목을 기재하고 전자결재 또는 서명을 완료하십시오.",
                    sheet_no=sheet_no,
                    sheet_title=sheet_title,
                )
            )
    return notes


# --- 규칙 ② 증빙·절차 ---
_ATTACH_REF_RE = re.compile(r"별\s*첨\s*(\d+|[A-Za-z가-힣]+)|첨\s*부\s*(\d+|[A-Za-z가-힣]+)", re.IGNORECASE)


def _check_evidence(doc: ParsedDocument) -> list[dict[str, Any]]:
    text = doc.text
    doc_no, doc_title = _doc_sheet(doc)
    notes: list[dict[str, Any]] = []

    if re.search(r"별\s*첨|첨\s*부", text) and not doc.tables:
        notes.append(
            _note(
                category="증빙·절차",
                importance="중",
                defect="‘별첨/첨부’ 표기에 해당하는 자료 미확인",
                reason="조서에 ‘별첨’ 또는 ‘첨부’ 표기가 있으나, 실제 첨부된 표·자료가 확인되지 않습니다.",
                basis="회계감사기준 500호(감사증거); 회계감사기준 230호(감사문서)",
                to_be="언급된 별첨 자료를 조서에 실제로 첨부하고 참조 위치를 명시하십시오.",
                sheet_no=doc_no,
                sheet_title=doc_title,
            )
        )

    # 별첨 번호 vs 실제 시트/표 존재 대조
    refs = set()
    for m in _ATTACH_REF_RE.finditer(text):
        ref = (m.group(1) or m.group(2) or "").strip()
        if ref:
            refs.add(ref)
    if refs and doc.tables:
        sheet_labels = set()
        for t in doc.tables:
            for val in (str(t.attrs.get("source", "")), str(t.attrs.get("title", ""))):
                sheet_labels.add(val)
                for part in re.split(r"[\s_\-]", val):
                    if part:
                        sheet_labels.add(part)
        missing = [r for r in refs if not any(r in lbl for lbl in sheet_labels)]
        if missing:
            notes.append(
                _note(
                    category="증빙·절차",
                    importance="중",
                    defect=f"별첨 참조({', '.join(missing[:5])}) — 해당 시트·표 미확인",
                    reason="본문에 별첨·첨부 참조가 있으나, 파싱된 시트/표 이름과 대응되지 않습니다.",
                    basis="회계감사기준 230호(감사문서); 500호(감사증거)",
                    to_be="별첨 참조 번호와 실제 첨부 시트명을 일치시키거나 누락 자료를 첨부하십시오.",
                    sheet_no=doc_no,
                    sheet_title=doc_title,
                )
            )

    return notes


def _check_account_mapping(doc: ParsedDocument) -> list[dict[str, Any]]:
    """계정 매핑 실패(미분류) 시트 경고."""
    notes: list[dict[str, Any]] = []
    for t in doc.tables:
        title = str(t.attrs.get("title", "")).strip()
        source = str(t.attrs.get("source", "")).strip()
        stext = _sheet_text(t)
        if len(stext.strip()) < 40:
            continue
        acct = table_account(t)
        if acct:
            continue
        if title and any(k in title for k in ("목차", "표지", "Index", "Summary")):
            continue
        notes.append(
            _note(
                category="QC·품질관리",
                importance="중",
                defect=f"계정 매핑 실패 — 「{source or title or '시트'}」 미분류",
                reason="해당 시트의 주계정을 식별하지 못했습니다. 절차누락·중점항목 검증에서 누락될 수 있습니다.",
                basis="시스템 안내 — 계정과목 매핑",
                to_be="시트 제목 또는 상단에 계정과목명을 명확히 기재하십시오.",
                sheet_no=source or "-",
                sheet_title=title or "미분류",
            )
        )
    return notes


# --- 규칙 ②-2 계정별 필수 감사절차 누락 ---
# 계정과목별로, 표준조서·감리지적사례에서 강조되는 핵심 감사절차를 정의한다.
#   detect: 절차 수행 흔적(키워드) — 하나라도 있으면 '수행'으로 간주
#   synonyms: 해당 계정이 조서의 주제임을 판단하는 계정명
PROCEDURE_RULES: list[dict[str, Any]] = [
    {
        "account": "매출채권",
        "synonyms": ["매출채권", "외상매출금", "받을어음"],
        "procs": [
            {
                "name": "채권 외부조회(조회서 발송·회신)",
                "detect": ["조회", "확인서", "잔액확인", "confirmation", "조회서"],
                "basis": "회계감사기준 505호(외부조회); 501호(특정 항목의 감사증거)",
                "to_be": "매출채권 잔액에 대한 조회서 발송·회신 현황(회신율·불일치 조정)을 조서에 기재하십시오.",
            },
            {
                "name": "대손충당금(기대신용손실) 평가",
                "detect": ["대손", "기대신용손실", "ECL", "손상", "연령분석", "aging", "회수가능"],
                "basis": "K-IFRS 1109호(금융상품); 일반기준 대손 관련 규정",
                "to_be": "대손충당금(기대신용손실) 산정 근거와 연령분석을 조서에 문서화하십시오.",
            },
        ],
    },
    {
        "account": "재고자산",
        "synonyms": ["재고자산", "상품", "제품", "원재료", "재공품"],
        "procs": [
            {
                "name": "재고자산 실사입회",
                "detect": ["실사", "입회", "재고조사"],
                "search_all_sheets": True,
                "basis": "회계감사기준 501호(재고자산 실사입회)",
                "to_be": "재고 실사 일시·대상·입회자 및 표본 대사 결과를 조서에 기재하십시오.",
            },
            {
                "name": "재고자산 평가(저가법·감모)",
                "detect": ["저가법", "평가손실", "감모", "순실현가능", "NRV", "진부화", "유효기간"],
                "basis": "K-IFRS 1002호(재고자산) 저가법 평가",
                "to_be": "저가법 평가(순실현가능가치)·감모·진부화 검토 내역을 조서에 문서화하십시오.",
            },
            {
                "name": "재고자산 Cut-off(기간귀속) 검토",
                "detect": [
                    "cut-off", "cutoff", "컷오프", "cut off", "기간귀속",
                    "결산일", "선적", "입고", "출고", "미착", "shipping",
                ],
                "search_all_sheets": True,
                "basis": "회계감사기준 501호; K-IFRS 1002호 재고자산 인식·기간귀속",
                "to_be": "결산일 전후 재고 입출고·선적 등 Cut-off 검토 내역을 조서에 기재하십시오.",
            },
        ],
    },
    {
        "account": "매입채무",
        "synonyms": ["매입채무", "외상매입금", "지급어음"],
        "procs": [
            {
                "name": "매입채무 외부조회·기간귀속",
                "detect": ["조회", "확인서", "잔액확인", "기간귀속", "cut-off", "cutoff", "미계상"],
                "basis": "회계감사기준 505호(외부조회); 완전성 관련 절차",
                "to_be": "매입채무 조회 또는 결산기 전후 매입 기간귀속(미계상 채무) 검토 내역을 기재하십시오.",
            },
        ],
    },
    {
        "account": "차입금",
        "synonyms": ["차입금", "사채", "장기차입", "단기차입"],
        "procs": [
            {
                "name": "차입금 조회·약정 확인",
                "detect": ["조회", "확인서", "잔액확인", "약정", "차입약정"],
                "basis": "회계감사기준 505호(외부조회)",
                "to_be": "차입금 잔액·이자율·약정사항에 대한 금융기관 조회 결과를 조서에 기재하십시오.",
            },
        ],
    },
    {
        "account": "현금및예금",
        "synonyms": ["현금및현금성자산", "예금", "현금성자산", "장단기금융상품"],
        "procs": [
            {
                "name": "예금 잔액조회(은행조회서)",
                "detect": [
                    "조회", "은행조회", "잔액조회", "확인서", "잔액확인",
                    "금융기관", "confirmation", " ref", "ref ",
                ],
                "search_all_sheets": True,
                "basis": "회계감사기준 505호(외부조회)",
                "to_be": "예금·금융상품 잔액에 대한 은행조회서 발송·회신 결과를 조서에 기재하십시오.",
            },
        ],
    },
    {
        "account": "매출",
        "synonyms": ["매출액", "수익", "매출원가"],
        "procs": [
            {
                "name": "수익 기간귀속(cut-off)",
                "detect": ["기간귀속", "cut-off", "cutoff", "마감", "귀속시점", "인식시점"],
                "search_all_sheets": True,
                "basis": "K-IFRS 1115호(고객과의 계약에서 생기는 수익) 기간귀속",
                "to_be": "결산기 전후 매출(수익) 기간귀속 검토(cut-off test) 내역을 조서에 기재하십시오.",
            },
        ],
    },
    {
        "account": "투자부동산",
        "synonyms": ["투자부동산"],
        "procs": [
            {
                "name": "공정가치 평가·방법",
                "detect": ["공정가치", "평가"],
                "detect_all": ["공정가치"],
                "basis": "K-IFRS 1040호(투자부동산)",
                "to_be": "공정가치 평가 방법·가정·감정보고서 등 근거를 조서에 기재하십시오.",
            },
            {
                "name": "임대수익·처분손익",
                "detect": ["임대", "처분"],
                "detect_all": [],
                "basis": "K-IFRS 1040호(투자부동산)",
                "to_be": "임대수익 인식·처분손익 계산 근거를 문서화하십시오.",
            },
        ],
    },
    {
        "account": "충당부채",
        "synonyms": ["충당부채", "제품보증충당"],
        "procs": [
            {
                "name": "충당부채 인식·측정",
                "detect": ["충당", "측정"],
                "detect_all": ["충당"],
                "basis": "K-IFRS 1037호(충당부채·우발부채)",
                "to_be": "충당부채 인식 요건·측정 근거를 조서에 기재하십시오.",
            },
        ],
    },
    {
        "account": "우발부채·약정",
        "synonyms": ["우발부채", "약정사항", "지급보증"],
        "procs": [
            {
                "name": "우발부채·약정 공시",
                "detect": ["우발", "약정", "공시"],
                "detect_all": [],
                "basis": "K-IFRS 1037호; K-IFRS 1001호(공시)",
                "to_be": "우발부채·약정사항 주석 공시 적정성을 검토·문서화하십시오.",
            },
        ],
    },
]


def _doc_evidence(doc: ParsedDocument) -> str:
    """문서 전체(모든 시트) 본문 — 타 조서에 수행된 절차 탐색용."""
    parts = [_sheet_text(t) for t in doc.tables]
    if not parts and doc.text:
        return doc.text
    return "\n".join(parts)


# Cut-off 등 타 조서에서 수행됐을 수 있는 절차 — notes_pipeline에서 후처리


def _subject_sheets(doc: ParsedDocument, account: str) -> list[tuple[str, str, str]]:
    """해당 계정이 '주(主)계정'인 시트들을 (조서번호, 계정과목, 시트본문)로 반환.

    각 시트의 주계정(_primary_account: 제목 → 본문 최초 등장 계정)이 해당 계정과
    일치하는 시트만 포함한다. 즉 그 시트가 실제로 그 계정을 다루는 조서여야 한다.
    합계·덤프에 다른 계정명이 섞여 있어도, 그 시트의 주제가 아니면 대상에서 제외된다.
    """
    result: list[tuple[str, str, str]] = []
    for t in doc.tables:
        title = str(t.attrs.get("title", ""))
        if table_account(t) == account:
            result.append((str(t.attrs.get("source", "")), title, _sheet_text(t)))
    return result


def _procedure_satisfied(evidence: str, proc: dict[str, Any]) -> bool:
    """절차 수행 흔적 — detect_all(AND) + detect(OR) 강화."""
    low = evidence.lower()
    detect_all = proc.get("detect_all") or []
    if detect_all and not all(k.lower() in low for k in detect_all):
        return False
    detect = proc.get("detect") or []
    return any(k.lower() in low for k in detect) if detect else bool(detect_all)


_TICK_MARK_RE = re.compile(r"(?:^|[\s,;|])[PV✓√○●](?:\s|$|[,.)])", re.M)
_CASH_CONFIRM_PROC: dict[str, Any] = {
    "detect": [
        "조회", "은행조회", "잔액조회", "확인서", "잔액확인",
        "금융기관", "confirmation", " ref", "ref ",
    ],
}


def _is_cash_sheet(t: pd.DataFrame) -> bool:
    if table_account(t) == "현금및예금":
        return True
    source = str(t.attrs.get("source", "")).strip()
    title = str(t.attrs.get("title", "")).strip()
    label = f"{source} {title}"
    code_part = source.split("/")[-1] if "/" in source else source
    if _account_from_code(code_part.split()[0] if code_part else "") == "현금및예금":
        return True
    return bool(re.search(r"\bA(?:\s|$|Lead|lead|리드)|현금|예금", label, re.I))


def doc_has_cash_confirmation_evidence(doc: ParsedDocument) -> bool:
    """A·A Lead 등 현금·예금 조서에 외부조회·tick·Ref 수행 흔적이 있는지."""
    doc_ev = _doc_evidence(doc).lower()
    if any(k in doc_ev for k in ("예금", "현금", "금융", "은행")):
        if _procedure_satisfied(doc_ev, _CASH_CONFIRM_PROC):
            return True
    for t in doc.tables:
        if not _is_cash_sheet(t):
            continue
        raw = _sheet_text(t)
        text = raw.lower()
        if _procedure_satisfied(text, _CASH_CONFIRM_PROC):
            return True
        if _TICK_MARK_RE.search(raw):
            return True
        if re.search(r"\bref\b", text, re.I):
            return True
    return False


def is_cash_external_confirm_note(note: dict[str, Any]) -> bool:
    """현금·예금 외부조회(은행조회서) 미수행 지적 노트인지."""
    hay = f"{note.get('defect', '')} {note.get('reason', '')}"
    if not re.search(
        r"외부\s?조회|조회서|금융기관|은행\s?조회|confirmation|잔액\s?(?:확인|조회)|예금\s?잔액",
        hay,
        re.I,
    ):
        return False
    acct = note_account(note)
    if acct == "현금및예금":
        return True
    sheet = f"{note.get('sheet_no', '')} {note.get('sheet_title', '')}"
    if re.search(r"(?:^|\s|/)A(?:\s|$|Lead|lead|리드)|현금|예금", sheet, re.I):
        return True
    return bool(re.search(r"현금|예금", hay, re.I))


_PHYSICAL_CASH_NA_RE = re.compile(
    r"(?:^|[\s,;|])(?:nan|n/a|n\.a\.|none|non|null|없음|해당\s?없|해당사항\s?없|"
    r"미보유|보유\s?없|실물\s?없|not\s?applicable)(?:[\s,;|]|$)",
    re.I,
)
_PHYSICAL_CASH_LABEL_RE = re.compile(
    r"^(?:\s*)?(?:현금(?![및성흐])|시재(?:\s?현금)?|자금(?:\s?시재)?|금고(?:\s?현금)?)\s*$",
    re.I,
)
_PHYSICAL_CASH_COUNT_RE = re.compile(
    r"현금\s?실사|시재\s?실사|실물\s?(?:현금|확인|검사)|현금\s?(?:입회|대사|확인)|"
    r"petty\s?cash|cash\s?count|physical\s?cash",
    re.I,
)
_DEPOSIT_LABEL_RE = re.compile(r"예금|금융상품|현금성|정기예금|보통예금|당좌", re.I)


@dataclass
class PhysicalCashInfo:
    """실물 현금(시재) 보유·실사 여부."""

    amount: float | None = None
    explicitly_none: bool = False
    has_count_evidence: bool = False


def _cell_is_na_marker(val: Any) -> bool:
    s = str(val).strip().lower()
    if not s or s in ("-", "—", "–"):
        return True
    if s == "nan":
        return True
    return bool(_PHYSICAL_CASH_NA_RE.search(s))


def _physical_cash_amount_from_text(text: str) -> tuple[float | None, bool]:
    """본문에서 실물 현금 행의 금액·n/a 표기 추출."""
    explicitly_none = False
    amount: float | None = None
    for line in text.splitlines():
        if not _PHYSICAL_CASH_LABEL_RE.search(line.split("\t")[0].strip() if "\t" in line else line[:20]):
            if not re.search(r"^(?:\s*)?(?:현금(?![및성흐])|시재)", line.strip(), re.I):
                continue
        if _DEPOSIT_LABEL_RE.search(line) and not re.search(r"^(?:\s*)?현금(?![및성흐])", line, re.I):
            continue
        if _cell_is_na_marker(line) or re.search(
            r"(?:nan|n/a|none|non|해당\s?없|해당사항\s?없|미보유|실물\s?없)", line, re.I
        ):
            explicitly_none = True
            continue
        for m in re.finditer(r"([\d,]+(?:\.\d+)?)", line):
            n = to_number(m.group(1))
            if n is not None:
                amount = max(amount or 0.0, abs(n))
    return amount, explicitly_none


def analyze_physical_cash(doc: ParsedDocument) -> PhysicalCashInfo:
    """조서 전체에서 실물 현금(시재) 보유액·n/a 표기·실사 흔적 분석."""
    info = PhysicalCashInfo()
    max_amt: float | None = None
    doc_ev = _doc_evidence(doc).lower()

    if _PHYSICAL_CASH_COUNT_RE.search(doc_ev):
        info.has_count_evidence = True

    for t in doc.tables:
        acct = table_account(t)
        source = str(t.attrs.get("source", "")).lower()
        title = str(t.attrs.get("title", "")).lower()
        is_cash_ctx = (
            acct == "현금및예금"
            or _is_cash_sheet(t)
            or "주석" in source
            or "주석" in title
        )
        if not is_cash_ctx:
            continue
        text = _sheet_text(t)
        if _PHYSICAL_CASH_COUNT_RE.search(text):
            info.has_count_evidence = True
        amt, na = _physical_cash_amount_from_text(text)
        if na:
            info.explicitly_none = True
        if amt is not None:
            max_amt = max(max_amt or 0.0, amt)

        # 표 형태: 첫 열 라벨이 '현금'·'시재'인 행
        if t.shape[1] >= 2:
            for _, row in t.iterrows():
                label = str(row.iloc[0]).strip()
                if not _PHYSICAL_CASH_LABEL_RE.match(label):
                    continue
                for val in row.iloc[1:]:
                    if _cell_is_na_marker(val):
                        info.explicitly_none = True
                        continue
                    n = to_number(val)
                    if n is not None:
                        max_amt = max(max_amt or 0.0, abs(n))

    info.amount = max_amt
    if info.amount is not None and info.amount <= 0:
        info.explicitly_none = True
    return info


def physical_cash_small_threshold(mat: Materiality | None = None) -> float:
    """소액 현금 기준 — 이하이면 실사 누락도 경미 지적으로만."""
    base = 1_000_000.0
    if mat and mat.ctt and mat.ctt > 0:
        return max(base, mat.ctt)
    if mat and mat.effective and mat.effective > 0:
        return max(base, mat.effective * 0.001)
    return base


def physical_cash_significant_threshold(mat: Materiality | None = None) -> float:
    """실물 현금 실사 지적 대상 — 이상일 때만 본격 지적."""
    small = physical_cash_small_threshold(mat)
    if mat and mat.effective and mat.effective > 0:
        return max(small * 3, mat.effective * 0.005, 3_000_000.0)
    return max(small * 3, 3_000_000.0)


def is_cash_physical_disclosure_note(note: dict[str, Any]) -> bool:
    """실물 현금·nan/n/a 공시·현금실사 관련 지적 노트."""
    hay = f"{note.get('defect', '')} {note.get('reason', '')} {note.get('to_be', '')}"
    acct = note_account(note)
    sheet = f"{note.get('sheet_no', '')} {note.get('sheet_title', '')}"
    cash_ctx = (
        acct == "현금및예금"
        or re.search(r"현금|시재|예금", hay, re.I)
        or re.search(r"주석|현금|예금", sheet, re.I)
    )
    if not cash_ctx:
        return False
    if re.search(
        r"nan|n/a|실물|시재|현금\s?실사|physical\s?cash|petty\s?cash|"
        r"세부(?:내역|항목).{0,20}(?:불완전|미기재|없)",
        hay,
        re.I,
    ):
        return True
    if re.search(r"공시.{0,24}(?:불완전|미흡|미기재)", hay) and re.search(
        r"현금|예금|nan", hay, re.I
    ):
        return True
    return False


_CROSS_SHEET_DISCLOSURE_GAP_RE = re.compile(r"공시|주석|disclosure|notes\s?to", re.I)
_CROSS_SHEET_DISCLOSURE_MISSING_RE = re.compile(
    r"미비|미흡|부족|미기재|누락|불완전|없(?:음|다|음)|포함(?:되)?지\s?않|"
    r"구체적(?:인)?.*(?:부족|없|미)|검토.*(?:미비|부족|미흡)|"
    r"내용.*(?:없|미|부족)|기재.*(?:없|미|부족)|평가.*(?:부족|미흡)",
    re.I,
)
_DISCLOSURE_CONTENT_KW = (
    "공시", "주석", "금액", "장부", "기말", "당기", "정책", "분류",
    "만기", "이자", "제약", "담보", "적격", "활성", "검토", "대사",
    "적정", "확인", "기재", "요약",
)


def is_cross_sheet_disclosure_note(note: dict[str, Any]) -> bool:
    """당해 시트에 주석·공시가 없다는 지적 — 타 Lead·주석 조서에 있을 수 있음."""
    hay = f"{note.get('defect', '')} {note.get('reason', '')} {note.get('to_be', '')}"
    if not _CROSS_SHEET_DISCLOSURE_GAP_RE.search(hay):
        return False
    return bool(_CROSS_SHEET_DISCLOSURE_MISSING_RE.search(hay))


def disclosure_note_account(note: dict[str, Any]) -> str | None:
    """공시 지적 노트의 대상 계정."""
    acct = note_account(note)
    if acct:
        return acct
    hay = f"{note.get('defect', '')} {note.get('reason', '')} {note.get('account', '')}"
    for name, syns in ACCOUNT_TAXONOMY:
        for s in (name, *syns):
            if _synonym_pos(hay, s) is not None:
                return name
    return None


def _normalize_sheet_key(label: str) -> str:
    part = str(label or "").split("/")[-1].strip().lower()
    return re.sub(r"\s+", "", part)


def _sheet_keys_match(exclude: str, label: str) -> bool:
    """exclude_sheet와 label이 동일 조서인지 — 'a'⊂'lead' 등 부분문자열 오탐 방지."""
    if not exclude:
        return False
    ex = _normalize_sheet_key(exclude)
    key = _normalize_sheet_key(label)
    if not ex or not key:
        return False
    if ex == key:
        return True
    ex_code = re.match(r"^([a-z]+\d*)", ex)
    key_code = re.match(r"^([a-z]+\d*)", key)
    if ex_code and key_code and ex_code.group(1) == key_code.group(1):
        return True
    return key.startswith(ex + "lead") or key.startswith(ex + "리드") or ex.startswith(key)


def _account_synonyms(account: str) -> tuple[str, ...]:
    for name, syns in ACCOUNT_TAXONOMY:
        if name == account:
            return (name, *syns)
    return (account,)


def doc_has_account_disclosure_elsewhere(
    doc: ParsedDocument,
    account: str | None,
    *,
    exclude_sheet: str = "",
) -> bool:
    """Lead·주석·타 조서에 해당 계정 공시·검토 내용이 있는지."""
    if not account:
        return False
    syns = _account_synonyms(account)

    for t in doc.tables:
        label = f"{t.attrs.get('source', '')} {t.attrs.get('title', '')}".strip()
        if _sheet_keys_match(exclude_sheet, label):
            continue
        text = _sheet_text(t)
        if not any(_synonym_pos(text, s) is not None for s in syns):
            continue
        is_lead_fn = _is_note_or_lead_table(t)
        if _extract_account_amount(t, account) is not None:
            return True
        kw_hits = sum(
            1 for k in _DISCLOSURE_CONTENT_KW
            if k in text.lower() or (len(k) > 2 and k in text)
        )
        if is_lead_fn and (kw_hits >= 2 or re.search(r"[\d,]{4,}", text)):
            return True
        if re.search(r"주석|공시|notes to|disclosure", f"{label} {text[:800]}", re.I):
            if kw_hits >= 1 or re.search(r"[\d,]{4,}", text):
                return True
    return False


def _check_materiality_missing(doc: ParsedDocument, mat: Materiality) -> list[dict[str, Any]]:
    if mat.effective:
        return []
    sheet_no, sheet_title = _doc_sheet(doc)
    return [
        _note(
            category="중요성",
            importance="상",
            defect="중요성 금액(TE/PM) 산정·문서화 미확인",
            reason="계획단계 조서에서 중요성 금액·수행중요성을 확인하지 못했습니다. "
            "KSA 320에 따라 중요성 산정 근거 문서화가 필수입니다.",
            basis="회계감사기준 320호(감사에서의 중요성)",
            to_be="중요성 금액·수행중요성·명백히 사소한 금액을 계획조서에 기재하십시오.",
            sheet_no=sheet_no,
            sheet_title=sheet_title,
        )
    ]


# KB 검색 결과 중 파일명·버전 표기만 있고 절차 의미가 없는 항목 식별
_FILELIKE_NAME_RE = re.compile(r"^[\d_./\s-]+(v\d|fy\s?\d{2,4}|\.xls)", re.I)


def _check_procedures(
    doc: ParsedDocument,
    mat: Materiality,
    is_listed: bool,
) -> list[dict[str, Any]]:
    """계정별 표준 감사절차 — Hanul DB 우선, 하드코딩 폴백."""
    notes: list[dict[str, Any]] = []
    kb_ready = kb.is_ready()
    use_fallback = app_config.fallback_to_hardcoded_procedures()
    doc_evidence = _doc_evidence(doc).lower()

    accounts_seen: set[str] = set()
    for name, _syns in ACCOUNT_TAXONOMY:
        sheets = _subject_sheets(doc, name)
        if not sheets:
            continue
        accounts_seen.add(name)
        sheet_no, title, _ = sheets[0]
        sheet_title = title or name
        evidence = "\n".join(s[2] for s in sheets).lower()

        specs: list[dict[str, Any]] = []
        if kb_ready:
            syns = _account_synonyms(name)
            for sp in kb.get_standard_procedures(name, is_listed=is_listed, synonyms=syns):
                if _FILELIKE_NAME_RE.match(sp.name.strip()):
                    continue
                detect = list(sp.detect_any)
                if sp.name and sp.name not in detect and len(sp.name) <= 20:
                    detect = [sp.name[:12], *detect]
                specs.append(
                    {
                        "name": sp.name,
                        "detect": detect[:6],
                        "detect_all": list(sp.detect_all),
                        "basis": sp.basis,
                        "to_be": sp.to_be,
                        "kb_snippet": sp.snippet,
                    }
                )

        rule = next((r for r in PROCEDURE_RULES if r["account"] == name), None)
        if rule:
            for p in rule["procs"]:
                if not any(p["name"] == s["name"] for s in specs):
                    specs.append(p)

        if not specs:
            # 표준절차 프로그램 미매칭은 시스템 진단 사항이므로 리뷰노트를 만들지 않는다.
            continue

        for proc in specs:
            ev = doc_evidence if proc.get("search_all_sheets") else evidence
            if _procedure_satisfied(ev, proc):
                continue
            imp = "상"  # 필수 절차 누락은 중요 흠결로 유지 (계정 잔액 기반 하향은 향후)
            reason_extra = ""
            if proc.get("search_all_sheets"):
                reason_extra = " (전체 조서를 검색했으나 해당 절차 흔적을 확인하지 못했습니다.)"
            kb_hint = ""
            if proc.get("kb_snippet"):
                snip = str(proc["kb_snippet"])[:100].replace("\n", " ")
                kb_hint = f" Hanul DB 표준·실무조서 참고: {snip}…"
            notes.append(
                _note(
                    category="절차누락",
                    importance=imp,
                    defect=f"[중요 절차 누락] {name} — {proc['name']} 흔적 미확인",
                    reason=(
                        f"「{name}」 조서에 {proc['name']} 수행 내역이 확인되지 않습니다."
                        f"{reason_extra} "
                        "표준조서·실무조서상 필수 절차이며, 누락 시 감리지적 대상이 될 수 있습니다."
                        f"{kb_hint}"
                    ),
                    basis=proc.get("basis", "회계감사기준"),
                    to_be=proc.get("to_be", "해당 절차를 수행·문서화하십시오."),
                    sheet_no=sheet_no,
                    sheet_title=sheet_title,
                )
            )

    # Hanul DB 미색인 등 시스템 상태는 리뷰노트가 아닌 화면 상태 칩으로만 안내한다.
    _ = (kb_ready, use_fallback)
    return notes


# --- 규칙 ③ 계산검증 ---
def _check_calculations(doc: ParsedDocument, mat: Materiality | None = None) -> list[dict[str, Any]]:
    notes: list[dict[str, Any]] = []
    m = mat or Materiality()
    for t_idx, table in enumerate(doc.tables, start=1):
        notes += _verify_table_totals(table, t_idx, m)
    return notes


def _check_footnotes(
    doc: ParsedDocument, mat: Materiality, *, include_minor: bool = False
) -> list[dict[str, Any]]:
    """주석 시트 필수기재·수치 정합성 점검."""
    notes: list[dict[str, Any]] = []
    footnote_tables = [
        t
        for t in doc.tables
        if "주석" in str(t.attrs.get("source", "")) or "주석" in str(t.attrs.get("title", ""))
    ]
    if include_minor and not footnote_tables and re.search(r"주석|Notes to", doc.text, re.IGNORECASE):
        sheet_no, sheet_title = _doc_sheet(doc)
        notes.append(
            _note(
                category="주석검증",
                importance="중",
                defect="주석 본문은 있으나 주석 표/시트 추출 미확인",
                reason="조서에 주석 관련 텍스트가 있으나 구조화된 주석 표를 파싱하지 못했습니다.",
                basis="K-IFRS 1001호(공시)",
                to_be="주석 Lead·주석 본문이 조서에 포함되어 있는지, PDF/Excel 형식을 확인하십시오.",
                sheet_no=sheet_no,
                sheet_title=sheet_title or "주석",
            )
        )

    for t_idx, table in enumerate(footnote_tables, start=1):
        notes += _verify_table_totals(table, t_idx, mat, force_footnote=True)
        # 주석 내 필수 키워드(공정가치, 충당, 우발 등) 존재 여부
        stext = _sheet_text(table).lower()
        title = str(table.attrs.get("title", ""))
        acct = table_account(table) or title
        required_kw = {
            "투자부동산": ("공정가치", "평가"),
            "충당부채": ("충당",),
            "우발부채·약정": ("우발", "약정"),
        }
        for key, kws in required_kw.items():
            if key in acct or key in title:
                if not all(k in stext for k in kws):
                    notes.append(
                        _note(
                            category="주석검증",
                            importance="상",
                            defect=f"주석 필수 공시항목 미확인 — {key}",
                            reason=f"주석 시트에서 {key} 관련 필수 공시 키워드({', '.join(kws)})가 확인되지 않습니다.",
                            basis="K-IFRS 1001호(공시); 금감원 중점심사 회계이슈",
                            to_be=f"{key} 주석 공시(금액·정책·위험)를 보완·문서화하십시오.",
                            sheet_no=str(table.attrs.get("source", "")),
                            sheet_title=title or key,
                        )
                    )
    return notes


# ── 주석·Lead 통합시트 검토 ─────────────────────────────────
# 주석사항은 계정별 통합(주석/Lead) 시트에서 집중 검토한다.
#   ① 각 조서 레퍼런스(A·C·E 등 조서번호)가 기재되어 있고 실제 시트와 대응되는지
#   ② 주석(Lead) 금액이 개별 조서 금액과 일치하게 추출되었는지
_NOTE_LEAD_RE = re.compile(r"주석|lead|리드|l/s", re.IGNORECASE)
_REF_TOKEN_RE = re.compile(r"(?<![A-Za-z가-힣])([A-Za-z]{1,3})[-.]?\d{0,3}(?![A-Za-z가-힣])")
_CURRENT_HDR_RE = re.compile(
    r"당\s*기|당해년|기\s*말|말\s*잔|CY|current|당기말|당기말",
    re.IGNORECASE,
)
_REF_TRUE_RE = re.compile(r"\b(?:True|TRUE)\b")


def _is_note_or_lead_table(t: pd.DataFrame) -> bool:
    label = f"{t.attrs.get('source', '')} {t.attrs.get('title', '')}"
    return bool(_NOTE_LEAD_RE.search(label))


def _sheet_code(label: str) -> str | None:
    return scr.parse_sheet_code(label)


def _sheet_ref_token(label: str) -> str:
    """시트명에서 조서번호 토큰 추출 (예: BB100)."""
    raw = str(label or "").strip().upper()
    if not raw:
        return ""
    tail = raw.split("/")[-1].strip()
    if not tail:
        return ""
    parts = tail.split()
    if not parts:
        return ""
    return parts[0].replace(",", ".")


def _row_text(table: pd.DataFrame, row_idx: int) -> str:
    return " ".join("" if c is None else str(c) for c in table.iloc[row_idx].tolist())


def _row_label_text(table: pd.DataFrame, row_idx: int, max_cols: int = 6) -> str:
    """행 계정·합계 라벨 — 병합·들여쓰기로 첫 열이 비어 있는 조서 대응."""
    parts: list[str] = []
    for ci in range(min(max_cols, table.shape[1])):
        v = table.iloc[row_idx, ci]
        s = ("" if v is None else str(v)).strip()
        if s and not re.match(r"^[\d,.\-]+$", s):
            parts.append(s)
    return " ".join(parts)


def _row_has_true_marker(table: pd.DataFrame, row_idx: int) -> bool:
    """조서 행에 Lead↔세부조서 대사 확인(True) 표시가 있는지."""
    return bool(_REF_TRUE_RE.search(_row_text(table, row_idx)))


def _amount_cols_current_year(table: pd.DataFrame) -> list[Any]:
    """당기(당해년도말) 금액 열 — 통상 오른쪽 숫자 열을 우선한다."""
    cached = table.attrs.get("_amount_cols_cy")
    if cached is not None:
        return cached
    if table.shape[1] < 2:
        table.attrs["_amount_cols_cy"] = []
        return []
    candidates: list[tuple[Any, bool, bool, int]] = []
    sample_n = min(40, len(table))
    for ci, col in enumerate(table.columns):
        if ci == 0:
            continue
        hdr = " ".join(
            str(table.iloc[r][col]) for r in range(min(5, len(table)))
        )
        is_prior = bool(_PRIOR_HDR_RE.search(hdr))
        is_current = bool(_CURRENT_HDR_RE.search(hdr))
        has_amount = any(
            (v := _pure_number(table.iloc[i][col])) is not None and abs(v) >= 1000
            for i in range(sample_n)
        )
        if has_amount:
            candidates.append((col, is_prior, is_current, ci))

    if not candidates:
        result = list(table.columns[1:])
    else:
        current_cols = [c for c, prior, cur, _ in candidates if cur and not prior]
        if current_cols:
            result = current_cols
        else:
            non_prior = [c for c, prior, _, _ in candidates if not prior]
            if non_prior:
                result = [non_prior[-1]]
            else:
                result = [candidates[-1][0]]
    table.attrs["_amount_cols_cy"] = result
    return result


def _row_amount_current_year(table: pd.DataFrame, row_idx: int) -> float | None:
    """행의 당기(오른쪽) 금액 — 복수 당기 열이면 가장 오른쪽 값."""
    amt, _ = _row_amount_with_ref(table, row_idx)
    return amt


def _excel_ref_for_cell(table: pd.DataFrame, row_idx: int, col: Any) -> str:
    """DataFrame 위치 → 엑셀 셀좌표 (예: D12)."""
    row_map = table.attrs.get("row_map") or []
    col_map = table.attrs.get("col_map") or []
    excel_row = int(row_map[row_idx]) if row_idx < len(row_map) else row_idx + 1
    cols = list(table.columns)
    col_idx = cols.index(col) if col in cols else -1
    excel_col = int(col_map[col_idx]) if 0 <= col_idx < len(col_map) else col_idx + 1
    return f"{_col_letter(excel_col)}{excel_row}"


def _row_amount_with_ref(table: pd.DataFrame, row_idx: int) -> tuple[float | None, str]:
    """행의 당기 금액과 해당 셀좌표."""
    cols = _amount_cols_current_year(table)
    if not cols:
        return None, ""
    for col in reversed(cols):
        v = _pure_number(table.iloc[row_idx][col])
        if v is not None and abs(v) >= 1000:
            return v, _excel_ref_for_cell(table, row_idx, col)
    return None, ""


def _table_label_column(table: pd.DataFrame) -> list[str]:
    """첫 열 라벨 목록 — 열이 없는 시트는 빈 리스트."""
    if table.shape[1] < 1:
        return []
    return [("" if v is None else str(v)) for v in table.iloc[:, 0].tolist()]


def _extract_account_amount_loc(
    table: pd.DataFrame, account: str
) -> tuple[float | None, str, str, int | None]:
    """통합(주석/Lead) 시트 — (금액, 셀좌표, 행라벨, 행인덱스)."""
    syns = _account_synonyms(account)
    labels = _table_label_column(table)
    best: float | None = None
    best_ref = ""
    best_lab = ""
    best_i: int | None = None
    for i, lab in enumerate(labels):
        if not any(_synonym_pos(lab, s) is not None for s in syns):
            continue
        v, ref = _row_amount_with_ref(table, i)
        if v is not None and (best is None or abs(v) > abs(best)):
            best, best_ref, best_lab, best_i = v, ref, _row_label_text(table, i) or lab.strip(), i
    return best, best_ref, best_lab, best_i


def _extract_sheet_total_loc(
    table: pd.DataFrame, account: str
) -> tuple[float | None, str, str, int | None]:
    """세부조서 — 합계/계정 행의 (금액, 셀좌표, 행라벨, 행인덱스)."""
    labels = _table_label_column(table)
    best: float | None = None
    best_ref = ""
    best_lab = ""
    best_i: int | None = None
    for i, lab in enumerate(labels):
        row_total = _TOTAL_RE.search(lab) or (account and account in lab)
        if not row_total:
            continue
        v, ref = _row_amount_with_ref(table, i)
        if v is not None and (best is None or abs(v) > abs(best)):
            best, best_ref, best_lab, best_i = v, ref, _row_label_text(table, i) or lab.strip(), i
    return best, best_ref, best_lab, best_i


def _sheet_short_code(source: str) -> str:
    code = scr.parse_sheet_code(source)
    if code:
        return code.upper()
    m = re.search(r"\b([A-Z]{1,3}\d{0,3})\b", (source or "").upper())
    return m.group(1) if m else (source or "").strip().upper()


def _collect_detail_amounts(
    doc: ParsedDocument, account: str, *, skip_lead: bool = True
) -> list[dict[str, Any]]:
    """동일 계정 세부조서별 금액·셀위치 목록 (포괄 대사용)."""
    syns = _account_synonyms(account)
    out: list[dict[str, Any]] = []
    seen: set[tuple[str, float]] = set()
    for at in doc.tables:
        if skip_lead and _is_note_or_lead_table(at):
            continue
        title = str(at.attrs.get("title", "")).strip()
        acct = table_account(at) or title
        if not acct or not any(_synonym_pos(acct, s) is not None for s in syns):
            if not any(_synonym_pos(title, s) is not None for s in syns):
                continue
        at_src = str(at.attrs.get("source", "")).strip()
        amt, ref, lab, _ = _extract_sheet_total_loc(at, account)
        if amt is None:
            amt, ref, lab, _ = _extract_account_amount_loc(at, account)
        if amt is None:
            continue
        key = (at_src, round(amt, 2))
        if key in seen:
            continue
        seen.add(key)
        out.append(
            {
                "source": at_src,
                "code": _sheet_short_code(at_src),
                "token": _sheet_ref_token(at_src),
                "amount": amt,
                "ref": ref,
                "label": lab,
            }
        )
    return out


def _format_amount_line(sheet: str, ref: str, amount: float, label: str = "") -> str:
    code = _sheet_short_code(sheet) or sheet
    loc = f"{code} {ref}" if ref else code
    lab = f" 「{label[:20]}」" if label else ""
    return f"{loc}{lab} {amount:,.0f}원"


def _extract_ref_tokens(text: str, *, own_code: str | None = None) -> set[str]:
    """본문에서 세부조서 레퍼런스 토큰(BB100 등) 추출."""
    refs: set[str] = set()
    for m in _REF_TOKEN_RE.finditer(text):
        token = m.group(0).upper().replace(".", "")
        code = m.group(1).upper()
        if not scr.parse_sheet_code(code):
            continue
        has_digits = any(ch.isdigit() for ch in m.group(0))
        if len(code) == 1 and not has_digits:
            continue
        if own_code and code == own_code:
            continue
        refs.add(token)
    return refs


def _lead_referenced_details(lead_table: pd.DataFrame, account: str) -> set[str]:
    """Lead 시트의 해당 계정 블록에서 레퍼런스된 세부조서 토큰."""
    syns = _account_synonyms(account)
    own = _sheet_code(str(lead_table.attrs.get("source", "")).strip())
    labels = _table_label_column(lead_table)
    acct_rows = [
        i for i, lab in enumerate(labels)
        if any(_synonym_pos(lab, s) is not None for s in syns)
    ]
    refs: set[str] = set()
    if acct_rows:
        block_start = max(0, min(acct_rows) - 1)
        block_end = min(len(labels), max(acct_rows) + 4)
        for i in range(block_start, block_end):
            refs.update(_extract_ref_tokens(_row_text(lead_table, i), own_code=own))
    else:
        for i, lab in enumerate(labels):
            refs.update(_extract_ref_tokens(_row_text(lead_table, i), own_code=own))
    return refs


def _lead_detail_cross_verified(
    lead_table: pd.DataFrame,
    detail_table: pd.DataFrame,
    account: str,
    lead_amt: float,
    detail_amt: float,
    detail_src: str,
    mat: Materiality,
) -> bool:
    """Lead↔세부조서 레퍼런스 대사가 조서상 True로 확인된 경우."""
    tol = max(mat.min_calc_threshold(), abs(lead_amt) * 0.005, abs(detail_amt) * 0.005)
    if _amounts_consistent(lead_amt, detail_amt, tol):
        return True

    detail_token = _sheet_ref_token(detail_src)
    detail_key = detail_token.replace(".", "")
    syns = _account_synonyms(account)
    lead_labels = _table_label_column(lead_table)
    detail_labels = _table_label_column(detail_table)

    for i in range(len(lead_table)):
        row_text = _row_text(lead_table, i)
        if not _REF_TRUE_RE.search(row_text):
            continue
        row_key = row_text.upper().replace(".", "")
        if detail_key and detail_key in row_key:
            return True

    for i in range(len(detail_table)):
        row_text = _row_text(detail_table, i)
        if not _REF_TRUE_RE.search(row_text):
            continue
        if re.search(r"lead|리드|주석", row_text, re.I):
            return True

    for i, lab in enumerate(lead_labels):
        if not any(_synonym_pos(lab, s) is not None for s in syns):
            continue
        if _row_has_true_marker(lead_table, i):
            return True

    for i, lab in enumerate(detail_labels):
        is_acct_row = any(_synonym_pos(lab, s) is not None for s in syns)
        is_total_row = bool(_TOTAL_RE.search(lab))
        if (is_acct_row or is_total_row) and _row_has_true_marker(detail_table, i):
            return True

    return False


def _extract_account_amount(table: pd.DataFrame, account: str) -> float | None:
    """통합(주석/Lead) 시트에서 해당 계정 행의 당기(오른쪽) 금액."""
    amt, _, _, _ = _extract_account_amount_loc(table, account)
    return amt


def _is_scale_multiple(a: float, b: float) -> bool:
    """단위 차이(천원·백만원 등)로 인한 배수 관계이면 True (오탐 방지).

    재무제표 레퍼런스는 통상 원 단위, 주석 요약은 백만원 단위로 기재하며
    이때 절사(내림)·반올림이 발생한다. 큰 값을 단위로 나눈 결과가 작은 값과
    최대 1단위(절사분)까지만 차이 나면 동일 금액의 단위 차이로 본다.
    예: 26,912,345,678원 ↔ 26,912백만원 (절사) → 정상.
    """
    if not a or not b:
        return False
    big, small = (abs(a), abs(b)) if abs(a) >= abs(b) else (abs(b), abs(a))
    if big == small:
        return False
    for k in (100.0, 1000.0, 10000.0, 100000.0, 1000000.0, 100000000.0):
        scaled = big / k
        if scaled < 0.5:
            break
        # 절사·반올림 허용: 축소 표시값과 1단위 이내 차이면 동일 금액으로 판정
        if abs(scaled - small) <= 1.0 + small * 0.001:
            return True
    return False


def _amounts_consistent(base: float, v: float, tol: float) -> bool:
    """두 시트의 금액이 같은 값(허용오차 내) 또는 단위 차이(절사 포함)이면 True."""
    return abs(base - v) <= tol or _is_scale_multiple(base, v)


# 주석 금액 ↔ 조서 합계 단순 대사를 하지 않는 계정.
# 우발부채·약정은 '당사가 제공한 지급보증'(예: 차입 관련 수출입은행 앞 제공분)과
# '제공받은 지급보증'이 성격상 분리되어 주석과 조서의 대표 금액이 다른 것이 정상이다.
# 실무 검토 방식: 금융기관조회서에서 조회된 금액이 주석(조회서 요약 또는 CL조서
# 레퍼런스)에 정확히 기재되었는지를 확인한다.
_NO_AMOUNT_TIEOUT_ACCOUNTS = {"우발부채·약정"}
_BORROWING_SUB_ROWS_RE = re.compile(
    r"단기\s*차입|유동성\s*장기|장기\s*차입|"
    r"단기차입금|유동성장기부채|장기차입금|단기차입|장기차입|^(?:사채)\s*$",
    re.I,
)
_BORROWING_INTEREST_SHEET_RE = re.compile(
    r"이자(?:비용|율)?|평균|평잔|weighted|interest|이자\s*계산|이자\s*테스트",
    re.I,
)
_BORROWING_BALANCE_CTX_RE = re.compile(
    r"명세|조회|대사|잔액|confirmation|유동성\s*분류",
    re.I,
)
_BORROWING_SHEET_RE = re.compile(r"BB|DD|차입|사채", re.I)
_CONFIRM_REF_RE = re.compile(r"조회서|조회사항|조회결과|confirmation", re.IGNORECASE)


def _is_borrowing_interest_test_sheet(
    source: str, title: str, sheet_text: str = ""
) -> bool:
    """평균 이자·이자비용 테스트 조서 — 잔액 교차합 대상에서 제외."""
    token = _sheet_ref_token(source).upper()
    try:
        import guidelines_loader as gl

        roles = gl.load_sheet_roles_from_db()
        if roles.get(token) in ("차입금_이자테스트", "이자테스트", "분석조서"):
            return True
    except Exception:  # noqa: BLE001
        pass
    if token == "BB200":
        return True
    hay = f"{source} {title} {sheet_text}"
    # 이자·평균 테스트는 '평균잔액' 등 잔액 키워드가 있어도 잔액 대사 대상이 아님
    if _BORROWING_INTEREST_SHEET_RE.search(hay):
        return True
    if re.search(r"BB\s*200|/BB200", hay, re.I):
        return True
    if _BORROWING_BALANCE_CTX_RE.search(hay):
        return False
    return False


def _cross_tieout_account_key(table: pd.DataFrame, acct: str) -> str:
    """시트 간 교차합 비교용 계정 키 (차입금 계열 통합)."""
    source = str(table.attrs.get("source", "")).strip()
    title = str(table.attrs.get("title", "")).strip()
    hay = f"{source} {title}"
    if acct == "차입금" or _BORROWING_SHEET_RE.search(hay):
        return "차입금"
    return acct


def _extract_borrowing_sheet_total(table: pd.DataFrame) -> float | None:
    """차입금 조서 — 합계 행 우선, 없으면 유형별(단기·유동성·장기) 합산."""
    labels = [_row_label_text(table, i) for i in range(len(table))]

    total_candidates: list[float] = []
    for i, lab in enumerate(labels):
        if _TOTAL_RE.search(lab):
            v = _row_amount_current_year(table, i)
            if v is not None:
                total_candidates.append(v)
    if total_candidates:
        return max(total_candidates, key=abs)

    sub_sum = 0.0
    sub_found = 0
    for i, lab in enumerate(labels):
        if _TOTAL_RE.search(lab):
            continue
        if _BORROWING_SUB_ROWS_RE.search(lab):
            v = _row_amount_current_year(table, i)
            if v is not None and abs(v) > 0:
                sub_sum += v
                sub_found += 1
    if sub_found >= 2:
        return sub_sum
    if sub_found == 1:
        return sub_sum

    for i, lab in enumerate(labels):
        lab_s = lab.strip()
        if lab_s in ("차입금", "장단기차입금", "장단기 차입금"):
            v = _row_amount_current_year(table, i)
            if v is not None:
                return v
    return None


def _filter_borrowing_tieout_entries(
    entries: list[tuple[str, float]],
) -> list[tuple[str, float]]:
    """이자 테스트·소액(이자비용 등)이 잔액 교차합에 섞이지 않게 제거."""
    kept = [
        (s, v)
        for s, v in entries
        if not _is_borrowing_interest_test_sheet(s, "", "")
    ]
    if len(kept) < 2:
        return kept
    max_abs = max(abs(v) for _, v in kept)
    if max_abs <= 0:
        return kept
    return [(s, v) for s, v in kept if abs(v) >= max_abs * 0.05]


def _cross_tieout_sheet_amount(table: pd.DataFrame, acct: str) -> float | None:
    """시트 간 교차합용 대표 금액 (차입금은 총합·이자테스트 조서 제외)."""
    tie_key = _cross_tieout_account_key(table, acct)
    cache_attr = f"_tieout_amt_{tie_key}"
    if cache_attr in table.attrs:
        return table.attrs[cache_attr]
    source = str(table.attrs.get("source", "")).strip()
    title = str(table.attrs.get("title", "")).strip()
    if tie_key == "차입금":
        stext = _sheet_text(table)[:1200]
        if _is_borrowing_interest_test_sheet(source, title, stext):
            table.attrs[cache_attr] = None
            return None
        val = _extract_borrowing_sheet_total(table)
    else:
        val = _extract_sheet_total(table, acct)
    table.attrs[cache_attr] = val
    return val


def _amount_in_table(table: pd.DataFrame, amount: float) -> bool:
    """금액이 시트 내 어느 셀에든 (단위 차이 포함) 존재하는지 확인."""
    if not amount:
        return False
    tol = max(1.0, abs(amount) * 0.001)
    for col in table.columns:
        for v in table[col].tolist():
            n = _pure_number(v)
            if n is None:
                continue
            if abs(n - amount) <= tol or _is_scale_multiple(n, amount):
                return True
    return False


def _check_contingency_confirmations(doc: ParsedDocument) -> list[dict[str, Any]]:
    """우발부채·지급보증 — 금융기관조회서 금액 반영 확인 (금액 대사 대신).

    제공한/제공받은 지급보증은 성격이 달라 금액 대사가 성립하지 않으므로,
    조회서 참조(조회서 요약 시트 또는 CL조서의 조회서 레퍼런스)가 확인되지 않는
    경우에만 1건의 확인 노트를 생성한다.
    """
    cont_tables = [t for t in doc.tables if table_account(t) == "우발부채·약정"]
    if not cont_tables:
        return []
    # 우발부채 시트·주석(Lead) 시트 어디에서든 조회서 언급이 있으면 검토된 것으로 본다
    for t in doc.tables:
        if any(t is ct for ct in cont_tables) or _is_note_or_lead_table(t):
            hay = (
                f"{t.attrs.get('source', '')} {t.attrs.get('title', '')} {_sheet_text(t)}"
            )
            if _CONFIRM_REF_RE.search(hay):
                return []
    src = str(cont_tables[0].attrs.get("source", "")).strip()
    title = str(cont_tables[0].attrs.get("title", "")).strip()
    return [
        _note(
            category="주석검증",
            importance="중",
            defect="우발부채·지급보증 — 금융기관조회서 금액 반영 확인 필요",
            reason="우발부채·약정(지급보증) 조서에서 금융기관조회서 참조가 확인되지 않습니다. "
            "제공한 보증과 제공받은 보증은 성격이 다르므로, 조회서에서 조회된 금액이 "
            "주석에 정확히 기재되었는지 확인이 필요합니다.",
            basis="회계감사기준 505호(외부조회); K-IFRS 1037호(우발부채 공시)",
            to_be="금융기관조회서 요약(또는 CL조서의 조회서 레퍼런스)과 주석 기재 금액을 "
            "대사하고 그 결과를 문서화하십시오.",
            sheet_no=src or "CL",
            sheet_title=title or "우발부채·약정",
        )
    ]


def _check_note_references(doc: ParsedDocument, mat: Materiality) -> list[dict[str, Any]]:
    """주석·Lead 통합시트 집중 검토 — 조서 레퍼런스·주석 금액 추출 정합."""
    notes: list[dict[str, Any]] = []
    lead_tables = [t for t in doc.tables if _is_note_or_lead_table(t)]
    if not lead_tables:
        return notes

    for t in lead_tables:
        source = str(t.attrs.get("source", "")).strip()
        title = str(t.attrs.get("title", "")).strip()
        stext = _sheet_text(t)

        # ① 조서 레퍼런스 기재 확인
        #  - 단독 한 글자(P, V, F 등)는 tick mark(대사·확인 표시)이지 조서 레퍼런스가
        #    아니므로 레퍼런스로도, 지적 대상으로도 취급하지 않는다.
        #  - 참조된 조서가 업로드(통합문서)에 없는 경우는 다른 파일의 조서를
        #    참조한 것일 수 있으므로 지적하지 않는다.
        refs: set[str] = set()
        tick_marks = 0
        for m in _REF_TOKEN_RE.finditer(stext):
            code = m.group(1).upper()
            if not scr.parse_sheet_code(code):
                continue
            has_digits = any(ch.isdigit() for ch in m.group(0))
            if len(code) == 1 and not has_digits:
                tick_marks += 1  # 예: P, V — tick mark로 간주
                continue
            refs.add(code)
        own = _sheet_code(source)
        if own:
            refs.discard(own)

        if not refs and tick_marks == 0:
            notes.append(
                _note(
                    category="주석검증",
                    importance="중",
                    defect=f"조서 레퍼런스 미기재 — 「{source or title}」",
                    reason="주석(Lead) 통합시트에서 개별 조서 참조(조서번호)나 대사 표시(tick mark)가 "
                    "확인되지 않습니다. 주석 수치와 개별 조서 간 추적이 어려울 수 있습니다.",
                    basis="회계감사기준 230호(감사문서); 법인 표준조서 상호참조 원칙",
                    to_be="주석(Lead) 시트의 각 항목에 근거 조서번호(레퍼런스)를 기재하십시오.",
                    sheet_no=source,
                    sheet_title=title or "주석/Lead",
                )
            )

        # ② 주석(Lead) 금액 ↔ 개별 조서 금액 추출 정합 (포괄 대사·셀좌표 명시)
        tol_min = mat.min_calc_threshold()
        checked_pairs: set[tuple[str, str]] = set()
        lead_code = _sheet_short_code(source)
        for at in doc.tables:
            if at is t or _is_note_or_lead_table(at):
                continue
            acct = table_account(at)
            if not acct:
                continue
            if acct in _NO_AMOUNT_TIEOUT_ACCOUNTS:
                continue
            at_src = str(at.attrs.get("source", "")).strip()
            at_token = _sheet_ref_token(at_src)
            at_code = _sheet_short_code(at_src)
            lead_refs = _lead_referenced_details(t, acct)
            if lead_refs and at_token and at_token not in lead_refs:
                continue

            pair_key = (acct, at_token or at_src)
            if pair_key in checked_pairs:
                continue
            checked_pairs.add(pair_key)

            detail_total, detail_ref, detail_lab, _ = _extract_sheet_total_loc(at, acct)
            if detail_total is None:
                detail_total, detail_ref, detail_lab, _ = _extract_account_amount_loc(at, acct)
            if detail_total is None:
                continue
            lead_amt, lead_ref, lead_lab, _ = _extract_account_amount_loc(t, acct)
            if lead_amt is None:
                continue
            if _lead_detail_cross_verified(t, at, acct, lead_amt, detail_total, at_src, mat):
                continue
            diff = abs(lead_amt - detail_total)
            tol = max(tol_min, abs(lead_amt) * 0.005, abs(detail_total) * 0.005)
            if diff < tol:
                continue
            if _is_scale_multiple(lead_amt, detail_total):
                continue
            if _amount_in_table(at, lead_amt):
                continue

            all_details = _collect_detail_amounts(doc, acct)
            matching = [
                d
                for d in all_details
                if d["source"] != at_src
                and (
                    _amounts_consistent(lead_amt, d["amount"], tol)
                    or _is_scale_multiple(lead_amt, d["amount"])
                )
            ]

            lead_line = _format_amount_line(source, lead_ref, lead_amt, lead_lab)
            detail_line = _format_amount_line(at_src, detail_ref, detail_total, detail_lab)
            tieout = (
                f"주석 {lead_ref or lead_code} {lead_amt:,.0f}원 ↔ "
                f"{at_code} {detail_ref} {detail_total:,.0f}원 (차이 {diff:,.0f}원)"
            )
            location = f"{lead_code} {lead_ref} · {at_code} {detail_ref}"

            if matching:
                alt = matching[0]
                alt_line = _format_amount_line(alt["source"], alt["ref"], alt["amount"], alt["label"])
                defect = f"주석-조서 레퍼런스·금액 정합 — {acct}"
                reason = (
                    f"주석(Lead) {lead_line} ↔ 참조 세부조서 {detail_line} (차이 {diff:,.0f}원). "
                    f"업로드 조서 전체 대사 시 {alt_line} 과(와) 일치합니다. "
                    f"레퍼런스({at_code}) 기재 오류·대사 대상 조서 착오 여부를 확인하십시오."
                )
                tieout += f" · 일치조서 {alt['code']} {alt['ref']} {alt['amount']:,.0f}원"
                to_be = (
                    f"주석 {lead_ref or lead_code} 행의 레퍼런스를 {alt['code']} 등 "
                    f"실제 일치 조서로 수정하거나, {at_code} {detail_ref} 금액 차이 원인을 문서화하십시오."
                )
            else:
                defect = f"주석-조서 금액 불일치 — {acct}"
                reason = (
                    f"주석(Lead) {lead_line} ↔ 세부조서 {detail_line} — 차이 {diff:,.0f}원. "
                    f"업로드 조서 전체에서 주석 금액과 일치하는 세부조서를 찾지 못했습니다. "
                    f"주석 추출(전기이월·단위·범위)·레퍼런스를 포괄 확인하십시오."
                )
                to_be = (
                    f"주석 {lead_ref or lead_code} ↔ {at_code} {detail_ref} 간 "
                    f"{diff:,.0f}원 차이 원인을 대사·문서화하십시오."
                )

            importance = "중"
            eff = mat.effective
            if eff and diff >= eff * 0.15:
                importance = "상"
            note = _note(
                category="주석검증",
                importance=importance,
                defect=defect,
                reason=reason,
                basis="K-IFRS 1001호(공시); 회계감사기준 520호(분석적 절차)",
                to_be=to_be,
                sheet_no=source,
                sheet_title=title or "주석/Lead",
                location=location,
            )
            note["tieout_detail"] = tieout
            notes.append(note)
    return notes


def _extract_sheet_total(table: pd.DataFrame, account: str) -> float | None:
    """시트에서 합계/총계 행 또는 계정명 행의 당기(오른쪽) 금액 추출."""
    amt, _, _, _ = _extract_sheet_total_loc(table, account)
    return amt


def _check_cross_fs_totals(doc: ParsedDocument, mat: Materiality) -> list[dict[str, Any]]:
    """Lead·FS 등 시트 간 동일 계정 금액 교차 대조 (기본)."""
    notes: list[dict[str, Any]] = []
    by_acct: dict[str, list[dict[str, Any]]] = {}
    for t in doc.tables:
        if _is_note_or_lead_table(t):
            continue
        title = str(t.attrs.get("title", "")).strip()
        source = str(t.attrs.get("source", "")).strip()
        acct = table_account(t) or title
        if not acct or len(acct) < 2:
            continue
        tie_key = _cross_tieout_account_key(t, acct)
        if tie_key in _NO_AMOUNT_TIEOUT_ACCOUNTS:
            continue
        val = _cross_tieout_sheet_amount(t, acct)
        if val is None:
            continue
        amt, ref, lab, _ = _extract_sheet_total_loc(t, acct)
        if amt is None:
            amt, ref, lab, _ = _extract_account_amount_loc(t, acct)
        by_acct.setdefault(tie_key, []).append(
            {
                "source": source or title,
                "code": _sheet_short_code(source),
                "amount": val,
                "ref": ref or "",
                "label": lab or "",
            }
        )

    tol_pct = 0.01
    min_mag = mat.min_calc_threshold()
    for acct, entries in by_acct.items():
        if acct == "차입금":
            filtered = _filter_borrowing_tieout_entries([(e["source"], e["amount"]) for e in entries])
            keep_src = {s for s, _ in filtered}
            entries = [e for e in entries if e["source"] in keep_src]
        if len(entries) < 2:
            continue
        vals = [e["amount"] for e in entries]
        if max(vals) - min(vals) < min_mag:
            continue
        if max(vals) > 0 and (max(vals) - min(vals)) / max(vals) < tol_pct:
            continue
        base = max(vals, key=abs)
        tol = max(min_mag, abs(base) * 0.005)
        if all(_amounts_consistent(base, v, tol) for v in vals):
            continue

        hi = max(entries, key=lambda e: abs(e["amount"]))
        lo = min(entries, key=lambda e: abs(e["amount"]))
        diff = abs(hi["amount"] - lo["amount"])
        loc_parts = [
            f"{e['code']} {e['ref']}" if e["ref"] else e["code"]
            for e in entries[:4]
        ]
        location = " · ".join(loc_parts)
        if len(entries) > 4:
            location += f" 외 {len(entries) - 4}건"
        sheet_lines = [
            _format_amount_line(e["source"], e["ref"], e["amount"], e["label"])
            for e in entries[:4]
        ]
        tieout = (
            f"{hi['code']} {hi['ref']} {hi['amount']:,.0f}원 ↔ "
            f"{lo['code']} {lo['ref']} {lo['amount']:,.0f}원 (차이 {diff:,.0f}원)"
        )
        note = _note(
            category="계산검증",
            importance="중",
            defect=f"교차합 불일치 — {acct}",
            reason=(
                f"동일 계정 「{acct}」 금액이 시트 간 일치하지 않습니다. "
                + " / ".join(sheet_lines[:3])
                + (f" 외 {len(entries) - 3}건" if len(entries) > 3 else "")
                + f" — 최대 차이 {diff:,.0f}원."
            ),
            basis="회계감사기준 520호(분석적 절차); Lead↔FS 대사",
            to_be=(
                f"{hi['code']} {hi['ref']} ↔ {lo['code']} {lo['ref']} 간 "
                f"{diff:,.0f}원 차이 원인을 대사·문서화하십시오."
            ),
            sheet_no=entries[0]["source"],
            sheet_title=acct,
            location=location,
        )
        note["tieout_detail"] = tieout
        notes.append(note)
    return notes


_PRIOR_HDR_RE = re.compile(r"전\s*기|prior|PY|전기", re.IGNORECASE)
_ANALYSIS_KW = ("분석", "변동", "증감", "원인", "전기대비", "increase", "decrease", "fluctuation")


def _check_prior_year_analysis(doc: ParsedDocument) -> list[dict[str, Any]]:
    """전기 금액 열 존재 시 변동 분석 문서화 여부."""
    notes: list[dict[str, Any]] = []
    for t in doc.tables:
        if t.shape[0] < 3 or t.shape[1] < 3:
            continue
        header = " ".join(str(t.iloc[r][c]) for r in range(min(3, len(t))) for c in t.columns)
        if not _PRIOR_HDR_RE.search(header):
            continue
        source = str(t.attrs.get("source", "")).strip()
        title = str(t.attrs.get("title", "")).strip()
        stext = _sheet_text(t)
        if any(k in stext for k in _ANALYSIS_KW):
            continue
        notes.append(
            _note(
                category="계산검증",
                importance="하",
                defect=f"전기 대비 분석 — 「{title or source}」 변동 원인 미확인",
                reason="전기·당기 비교 열이 있으나 변동 원인·분석적 검토 서술이 확인되지 않습니다.",
                basis="회계감사기준 520호(분석적 절차)",
                to_be="전기/당기 비교표, 증감 원인, 추가 절차 필요 시 수행 내역을 기재하십시오.",
                sheet_no=source or "-",
                sheet_title=title or "전기대비",
            )
        )
    return notes


_MIN_TOTAL_MAGNITUDE = 1000  # 비율·건수·연도 등 소액 값은 합계검증 제외
_MIN_REL_DIFF = 0.03  # 상대 차이 3% 미만이면 오탐 가능성으로 보고 생략

# 순수 숫자 셀만 허용 (예: '12,000', '(1,500)', '3.5%' O / '2024감사후', '2024-12-31' X)
_PURE_NUM_RE = re.compile(r"^[\s(]*-?[\d,]+(?:\.\d+)?\s*%?[\s)]*$")
# 합계검증 대상이 아닌 열 머리글 (비율·건수·코드 등)
_SKIP_COL_HDR = re.compile(
    r"%|비율|률|건수|번호|코드|연도|년도|page|페이지|비고|주석번호",
    re.IGNORECASE,
)
# 차감·환입 성격 행 — 합산 시 부호 반전 후보 (일반 '비용' 행은 제외)
_DEDUCT_ROW = re.compile(r"대손|충당|차감|환입|감소|타계정", re.IGNORECASE)


def _pure_number(v: Any) -> float | None:
    """텍스트가 섞이지 않은 순수 숫자 셀만 값으로 반환 (머리글·날짜 제외)."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return to_number(v)
    s = str(v).strip()
    if not s or not _PURE_NUM_RE.match(s):
        return None
    return to_number(s)


def _col_values(table: pd.DataFrame, block: list[int], col: Any, labels: list[str]) -> list[float]:
    """해당 열에 값이 있는 행만 모아 반환(열 지정 오류 방지). 차감 행은 부호 반전 후보."""
    out: list[float] = []
    for pos in block:
        v = _pure_number(table.iloc[pos][col])
        if v is None:
            continue
        if _DEDUCT_ROW.search(labels[pos]) and v > 0:
            out.append(-v)
        else:
            out.append(v)
    return out


def _sum_variants(table: pd.DataFrame, block: list[int], col: Any, labels: list[str]) -> list[float]:
    """합계 해석 후보: 단순합·차감반영·양수만·소계행만 등."""
    if len(block) < 2:
        return []

    raw = [_pure_number(table.iloc[p][col]) for p in block]
    present = [v for v in raw if v is not None]
    if len(present) < 2:
        return []

    variants: list[float] = [sum(present)]

    signed = _col_values(table, block, col, labels)
    if len(signed) >= 2:
        variants.append(sum(signed))

    positives = [v for v in present if v > 0]
    if len(positives) >= 2:
        variants.append(sum(positives))

    # 소계 행만 합산 (계층표: 세부 항목 대신 소계끼리 합치는 경우)
    sub_vals = [
        _pure_number(table.iloc[p][col])
        for p in block
        if _SUBTOTAL_RE.search(labels[p]) and _pure_number(table.iloc[p][col]) is not None
    ]
    if len(sub_vals) >= 1:
        variants.append(sum(sub_vals))
        if len(sub_vals) == 1:
            variants.append(sub_vals[0])

    # 직전 소계/단일 행이 곧 합계인 경우
    if present:
        variants.append(present[-1])
    if len(present) >= 2:
        variants.append(sum(present[-2:]))

    return variants


def _matches_total(total_val: float, candidates: list[float], tolerance: float) -> bool:
    """여러 합산 해석 중 하나라도 합계와 맞으면 통과."""
    for s in candidates:
        if abs(s - total_val) <= tolerance:
            return True
    return False


_SUBSET_POOL_MAX = 24  # 부분합 탐색 대상 행 수 상한 (meet-in-the-middle로 계산)


def _subset_sum_matches(total_val: float, pool: list[float], tolerance: float) -> bool:
    """위쪽 행들의 '어떤 부분집합'의 합이 합계와 일치하는지 검사.

    감사조서의 합계행은 직전 행 전체가 아니라 일부 행만 골라 합산한
    부분합인 경우가 많다(예: 특정 항목 제외, 이전 소계 참조 포함).
    이 경우 오류로 지적하면 안 되므로 부분집합 합 일치 여부를 확인한다.
    """
    vals = [v for v in pool if v is not None and abs(v) >= 1]
    if not vals:
        return False
    vals = vals[-_SUBSET_POOL_MAX:]
    half = len(vals) // 2
    a, b = vals[:half], vals[half:]

    def _sums(items: list[float]) -> list[float]:
        out = [0.0]
        for v in items:
            out += [s + v for s in out]
        return out

    sums_a = _sums(a)
    sums_b = sorted(_sums(b))
    for sa in sums_a:
        target = total_val - sa
        lo, hi = 0, len(sums_b) - 1
        while lo <= hi:
            mid = (lo + hi) // 2
            if sums_b[mid] < target - tolerance:
                lo = mid + 1
            elif sums_b[mid] > target + tolerance:
                hi = mid - 1
            else:
                # 공집합·전체합(이미 검사됨)이 아닌 실제 부분합만 인정
                if abs(sa) >= 1 or abs(sums_b[mid]) >= 1:
                    return True
                break
    return False


def _is_amount_column(table: pd.DataFrame, col: Any, header_pos: int | None, labels: list[str]) -> bool:
    """금액 열인지 판단 — 비율·건수 열은 합계검증 제외."""
    if header_pos is not None and 0 <= header_pos < len(table):
        hdr = str(table.iloc[header_pos][col]).strip()
        if hdr and _SKIP_COL_HDR.search(hdr):
            return False
    nums: list[float] = []
    pct = 0
    for i, label in enumerate(labels):
        v = _pure_number(table.iloc[i][col])
        if v is None:
            continue
        s = str(table.iloc[i][col]).strip()
        if s.endswith("%"):
            pct += 1
            continue
        nums.append(v)
    if not nums:
        return False
    if pct > len(nums) * 0.5:
        return False
    large = sum(1 for n in nums if abs(n) >= _MIN_TOTAL_MAGNITUDE)
    return large >= max(2, len(nums) * 0.3)


_CELL_REF_RE = re.compile(
    r"(?:'([^']+)'!)?(\$?)([A-Z]{1,3})(\$?)(\d+)",
    re.IGNORECASE,
)
_SUM_RANGE_RE = re.compile(r"SUM\(([^)]+)\)", re.IGNORECASE)
_INT_FN_RE = re.compile(r"INT\(([^)]+)\)", re.IGNORECASE)


def _excel_col_index(letter: str) -> int:
    from openpyxl.utils import column_index_from_string

    return column_index_from_string(letter.upper())


def _table_pos_for_excel(
    table: pd.DataFrame, excel_row: int, excel_col: int
) -> tuple[int, Any] | None:
    row_map = table.attrs.get("row_map") or []
    col_map = table.attrs.get("col_map") or []
    table_row = next((i for i, r in enumerate(row_map) if r == excel_row), None)
    table_col_idx = next((i for i, c in enumerate(col_map) if c == excel_col), None)
    if table_row is None or table_col_idx is None:
        return None
    cols = list(table.columns)
    if table_col_idx >= len(cols):
        return None
    return table_row, cols[table_col_idx]


def _excel_cell_amount(table: pd.DataFrame, col_letter: str, excel_row: int) -> float | None:
    pos = _table_pos_for_excel(table, excel_row, _excel_col_index(col_letter))
    if pos is None:
        return None
    tr, tc = pos
    return _pure_number(table.iloc[tr][tc])


def _expand_sum_range(range_spec: str, get_cell) -> list[float]:
    vals: list[float] = []
    spec = range_spec.strip().upper()
    if ":" in spec:
        start, end = spec.split(":", 1)
        m1 = _CELL_REF_RE.search(start.strip())
        m2 = _CELL_REF_RE.search(end.strip())
        if m1 and m2:
            c1, r1 = m1.group(3), int(m1.group(5))
            c2, r2 = m2.group(3), int(m2.group(5))
            if c1 == c2:
                for r in range(min(r1, r2), max(r1, r2) + 1):
                    v = get_cell(c1, r)
                    if v is not None:
                        vals.append(v)
    else:
        for part in spec.split(","):
            m = _CELL_REF_RE.search(part.strip())
            if m:
                v = get_cell(m.group(3), int(m.group(5)))
                if v is not None:
                    vals.append(v)
    return vals


def _eval_cell_formula(table: pd.DataFrame, formula: str) -> float | None:
    """엑셀 셀 산식을 표 내 참조값으로 해석·계산."""
    if not formula or not str(formula).startswith("="):
        return None
    expr = str(formula).lstrip("=").strip()

    def get_cell(col_letter: str, row: int) -> float | None:
        return _excel_cell_amount(table, col_letter, row)

    while True:
        m = _SUM_RANGE_RE.search(expr)
        if not m:
            break
        parts = _expand_sum_range(m.group(1), get_cell)
        repl = str(sum(parts)) if parts else "0"
        expr = expr[: m.start()] + repl + expr[m.end() :]

    expr = _INT_FN_RE.sub(r"(\1)", expr)

    def repl_ref(m: re.Match[str]) -> str:
        v = get_cell(m.group(3), int(m.group(5)))
        return str(v) if v is not None else "0"

    expr = _CELL_REF_RE.sub(repl_ref, expr)
    expr = expr.replace(" ", "")
    if not re.match(r"^[-+]?[\d.eE+\-*/().]+$", expr):
        return None
    try:
        result = eval(expr, {"__builtins__": {}}, {})  # noqa: S307 — 숫자·연산자만 허용
        return float(result)
    except Exception:
        return None


def _formula_validates_total(
    table: pd.DataFrame,
    pos: int,
    col: Any,
    total_val: float,
    tolerance: float,
    excel_row_fn,
) -> bool:
    """합계 셀 산식(차감·SUM·셀참조) 결과가 기재값과 일치하면 오류 아님."""
    formulas = table.attrs.get("cell_formulas") or {}
    col_map = table.attrs.get("col_map") or []
    try:
        col_idx = list(table.columns).index(col)
    except ValueError:
        col_idx = int(col) if isinstance(col, int) else 0
    excel_col = col_map[col_idx] if col_idx < len(col_map) else col_idx + 1
    excel_row = excel_row_fn(pos)
    formula = formulas.get((excel_row, excel_col))
    if not formula:
        return False
    computed = _eval_cell_formula(table, formula)
    if computed is None:
        return False
    return abs(computed - total_val) <= tolerance


def _difference_pair_matches(total_val: float, pool: list[float], tolerance: float) -> bool:
    """두 행 금액의 차(마이너스 산식)가 합계와 일치하면 정상."""
    vals = [v for v in pool if v is not None and abs(v) >= 1]
    for a in vals:
        for b in vals:
            if abs(a - b - total_val) <= tolerance or abs(b - a - total_val) <= tolerance:
                return True
    return False


_MAX_CALC_VERIFY_ROWS = 500


def _verify_table_totals(
    table: pd.DataFrame,
    t_idx: int,
    mat: Materiality | None = None,
    *,
    force_footnote: bool = False,
) -> list[dict[str, Any]]:
    """합계행 대비 구성행 합산 — 맥락(열·범위·차감)을 여러 방식으로 해석해 오탐을 최소화."""
    notes: list[dict[str, Any]] = []
    m = mat or Materiality()
    min_mag = m.min_calc_threshold()
    if table.shape[0] < 3 or table.shape[1] < 2:
        return notes

    sheet_no = table.attrs.get("source", f"표 {t_idx}")
    sheet_title = table.attrs.get("title", "")
    sheet = f"{sheet_no} ({sheet_title})" if sheet_title else sheet_no
    row_map = table.attrs.get("row_map")
    is_footnote = "주석" in sheet_no or "주석" in sheet_title
    labels = _table_label_column(table)

    def excel_row(pos: int) -> int:
        if row_map and 0 <= pos < len(row_map):
            return row_map[pos]
        return pos + 1

    def excel_cell(pos: int, col: Any) -> str:
        row = excel_row(pos)
        if isinstance(col, int):
            return f"{_col_letter(col + 1)}{row}"
        return f"{row}행"

    pure_count = table.apply(
        lambda row: sum(_pure_number(v) is not None for v in row[1:]), axis=1
    ).tolist()
    is_data = [
        pure_count[i] >= 1 and bool(labels[i].strip()) and to_number(labels[i]) is None
        for i in range(len(labels))
    ]

    def is_total(pos: int) -> bool:
        return bool(_TOTAL_RE.search(labels[pos]))

    def is_subtotal(pos: int) -> bool:
        return bool(_SUBTOTAL_RE.search(labels[pos]))

    def col_label(col: Any, header_pos: int | None) -> str:
        if header_pos is not None and 0 <= header_pos < len(table):
            label = str(table.iloc[header_pos][col]).strip()
            if label and label.lower() != "nan":
                return label
        return f"{col}열"

    n = len(table)
    seen_cols: set = set()
    if n > _MAX_CALC_VERIFY_ROWS:
        scan_positions = [i for i in range(n) if is_data[i] and is_total(i)]
    else:
        scan_positions = range(n)
    for pos in scan_positions:
        if not (is_data[pos] and is_total(pos)):
            continue

        region_start = 0
        for j in range(pos - 1, -1, -1):
            if is_total(j):
                region_start = j + 1
                break

        region_block = [
            i
            for i in range(region_start, pos)
            if is_data[i] and not is_total(i) and not is_subtotal(i)
        ]
        contig_block: list[int] = []
        j = pos - 1
        while j >= region_start and is_data[j] and not is_total(j):
            if not is_subtotal(j):
                contig_block.insert(0, j)
            j -= 1

        # 합계 직전 소계만 대상으로 하는 후보③
        subtotal_only: list[int] = []
        for j in range(pos - 1, region_start - 1, -1):
            if is_subtotal(j):
                subtotal_only.insert(0, j)
                break
            if is_total(j):
                break

        blocks = [b for b in (contig_block, region_block, subtotal_only) if len(b) >= 2]
        if not blocks:
            continue
        primary_block = contig_block if len(contig_block) >= 2 else blocks[0]
        header_pos = primary_block[0] - 1 if primary_block else None

        for col in table.columns[1:]:
            total_val = _pure_number(table.iloc[pos][col])
            if total_val is None or abs(total_val) < min_mag:
                continue
            if 1900 <= abs(total_val) <= 2100 and float(total_val).is_integer():
                continue
            if not _is_amount_column(table, col, header_pos, labels):
                continue

            tolerance = max(1.0, abs(total_val) * 0.005)

            # 산식·차감 우선 확인 — =D9-D10 등 마이너스 산식은 오류가 아님
            if _formula_validates_total(table, pos, col, total_val, tolerance, excel_row):
                continue
            col_pool = [
                v
                for p in range(n)
                if (v := _pure_number(table.iloc[p][col])) is not None
            ]
            subset_tol = max(1.0, abs(total_val) * 0.001)
            if _difference_pair_matches(total_val, col_pool, subset_tol):
                continue

            all_candidates: list[float] = []
            for block in blocks:
                all_candidates.extend(_sum_variants(table, block, col, labels))

            # 구성행 개별 값이 합계와 일치 (참조·소계)
            for block in blocks:
                for p in block:
                    v = _pure_number(table.iloc[p][col])
                    if v is not None:
                        all_candidates.append(v)

            if _matches_total(total_val, all_candidates, tolerance):
                continue
            if not all_candidates:
                continue

            best = min(all_candidates, key=lambda s: abs(s - total_val))
            diff = best - total_val
            if abs(total_val) > 0 and abs(diff) / abs(total_val) < _MIN_REL_DIFF:
                eff = m.effective
                if not (eff and abs(diff) >= eff * 0.05):
                    continue
            # 차이 ≈ 빠진/추가된 한 행 금액 → 합산 범위 지정 문제로 보고 생략
            if any(
                _pure_number(table.iloc[p][col]) is not None
                and abs(abs(_pure_number(table.iloc[p][col]) or 0) - abs(diff)) <= tolerance
                for p in region_block
            ):
                continue
            # 부분합 검사 — 합계가 위쪽 행들 중 '일부 행'의 합(부분합)이면 정상.
            # 라벨이 비었거나 이전 소계·합계 행이 참조된 경우까지 포함해 확인한다.
            pool = [
                v
                for p in range(pos)
                if (v := _pure_number(table.iloc[p][col])) is not None
            ]
            if _subset_sum_matches(total_val, pool, subset_tol):
                continue

            name = col_label(col, header_pos)
            key = (sheet, name, round(total_val, 2))
            if key in seen_cols:
                continue
            seen_cols.add(key)

            cell = excel_cell(pos, col)
            comp_from = excel_row(primary_block[0])
            comp_to = excel_row(primary_block[-1])
            location = f"합계셀 {cell} · 인접 구성행 {comp_from}~{comp_to}행"
            importance = "하"
            eff = m.effective
            if eff and abs(diff) >= eff * 0.05:
                importance = "중"
            if eff and abs(diff) >= eff * 0.15:
                importance = "상"
            mat_note = ""
            if eff and eff > 0:
                pct = abs(diff) / eff * 100
                mat_note = f" (차액 {abs(diff):,.0f}원, 중요성의 {pct:.1f}%)"
            summary = (
                f"〈{sheet}〉 {cell} ‘{name}’ 합계 재확인 권장 "
                f"(기재 {total_val:,.0f} / 인접행 합 {best:,.0f}, 차이 {diff:,.0f}){mat_note} — "
                "합계 수식의 합산 범위·열·차감(마이너스) 지정을 확인하십시오."
            )
            notes.append(
                _note(
                    category="계산검증",
                    importance=importance,
                    location=location,
                    summary=summary,
                    defect=f"합계 재확인 — {cell} ({name})",
                    reason=(
                        f"‘{sheet}’ {cell} 합계({total_val:,.0f})와 "
                        f"직상 인접 구성행 {comp_from}~{comp_to}행의 합({best:,.0f})이 "
                        f"{diff:,.0f} 차이 납니다. "
                        "감사조서에서는 합계 수식의 범위·열·차감 항목 지정에 따라 "
                        "일치하는 경우가 많으므로, 수식 참조를 우선 확인하십시오."
                    ),
                    basis="회계감사기준 520호(분석적 절차)",
                    to_be=(
                        "해당 합계 셀의 SUM(또는 소계 참조) 범위·열·차감(음수) 항목이 "
                        "올바른지 확인하고, 필요 시 조서를 수정하십시오."
                    ),
                    sheet_no=sheet_no,
                    sheet_title=sheet_title,
                )
            )
    return notes


# --- 보조 함수 ---
def _note(
    *,
    category: str,
    defect: str,
    reason: str,
    basis: str,
    to_be: str,
    sheet_no: str = "",
    sheet_title: str = "",
    location: str = "",
    summary: str = "",
    importance: str | None = None,
    is_focus_related: bool = False,
) -> dict[str, Any]:
    sheet_no = (sheet_no or "").strip()
    sheet_title = (sheet_title or "").strip()
    if sheet_no and sheet_title:
        label = f"{sheet_no} ({sheet_title})"
    else:
        label = sheet_no or sheet_title or "조서 본문"
    imp = importance or _IMPORTANCE_BY_CATEGORY.get(category, "중")
    if not summary:  # 경미 사항용 한 줄 요약 기본값
        loc = f" {location}" if location else ""
        summary = f"〈{sheet_no or label}{loc}〉 {defect} — {to_be}"
    return {
        "id": "",
        "importance": imp,
        "category": category,
        "defect": defect,
        "reason": reason,
        "basis": basis,
        "to_be": to_be,
        "sheet_no": sheet_no or "-",
        "sheet_title": sheet_title,
        "sheet": label,
        "location": location,
        "summary": summary,
        "workpaper_ref": sheet_no or label,
        "is_focus_related": is_focus_related,
        "source": "rule",
    }


def _assign_ids(notes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    for i, note in enumerate(notes, start=1):
        note["id"] = f"RN-{i:03d}"
    return notes


def _search(pattern: str, text: str) -> str | None:
    m = re.search(pattern, text)
    if not m:
        return None
    return m.group(1).strip() if m.groups() else m.group(0).strip()
